from __future__ import annotations

"""Parecer de evolução contratual em Word (.docx), editável.

A abertura e o fechamento vêm do próprio modelo corporativo
(templates/modelo_manifestacao_subsidios.docx): faixa institucional, quadros
01 a 03, rótulo de classificação, rodapé e quadro "Responsável pela
informação". Os campos conhecidos pelo Motor são preenchidos; os demais
(por exemplo, Mutuário(s) e Matrícula(s), bloqueados pela LGPD no modo web)
ficam disponíveis para preenchimento manual no Word.

O conteúdo técnico é o mesmo do PDF: o story do ReportLab é convertido em
parágrafos e tabelas nativos do Word, com a mesma fonte, cores e tamanhos.
Ícones decorativos vetoriais não são levados para o Word.
"""

from datetime import date
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.graphics.shapes import Drawing
from reportlab.lib.colors import Color, toColor
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import (
    CondPageBreak, HRFlowable, Image, KeepTogether, PageBreak, Paragraph, Spacer, Table,
)

from . import pdf_funcef_template as funcef
from . import pdf_reports as base
from . import pdf_reports_v13 as v13
from .pdf_funcef_template import ManifestationHeader

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / "templates" / "modelo_manifestacao_subsidios.docx"
_TEMPLATE_DATE = "21/09/2026"
_FONT = "Arial"
_CONTENT_WIDTH_TW = 10602  # largura útil do modelo (A4, margens de 11,5 mm)
_VALUE_COLOR = "17212B"
_FOOTER_COLOR = "8A949E"
_ALIGN = {0: "left", 1: "center", 2: "right", 4: "both"}
_CELL_ALIGN = {"LEFT": "left", "CENTER": "center", "CENTRE": "center", "RIGHT": "right", "DECIMAL": "right"}
_VALIGN = {"TOP": "top", "MIDDLE": "center", "BOTTOM": "bottom"}


def docx_filename(modality_name: str, contract_number: str) -> str:
    return base.pdf_filename(modality_name, contract_number)[:-4] + ".docx"


# ---------------------------------------------------------------------------
# Utilidades de XML
# ---------------------------------------------------------------------------

def _el(tag: str, parent=None, **attrs):
    element = OxmlElement(tag)
    for key, value in attrs.items():
        element.set(qn(f"w:{key}"), str(value))
    if parent is not None:
        parent.append(element)
    return element


def _hex(color) -> str | None:
    if color is None:
        return None
    if not isinstance(color, Color):
        try:
            color = toColor(color)
        except Exception:
            return None
    return "%02X%02X%02X" % tuple(int(round(c * 255)) for c in (color.red, color.green, color.blue))


def _tw(points: float) -> int:
    return int(round(points * 20))


_KEEP_TABLE_ROWS = 14


def _keep_paragraphs(element) -> None:
    paragraphs = [element] if element.tag == qn("w:p") else list(element.iter(qn("w:p")))
    for paragraph in paragraphs:
        ppr = paragraph.find(qn("w:pPr"))
        if ppr is None:
            ppr = _el("w:pPr")
            paragraph.insert(0, ppr)
        if ppr.find(qn("w:keepNext")) is None:
            ppr.insert(0, _el("w:keepNext"))


def _text_of(element) -> str:
    return "".join(t.text or "" for t in element.iter(qn("w:t"))).strip()


# ---------------------------------------------------------------------------
# Texto com marcação do ReportLab → runs do Word
# ---------------------------------------------------------------------------

class _Markup(HTMLParser):
    def __init__(self, base_format: dict):
        super().__init__(convert_charrefs=True)
        self.stack = [base_format]
        self.parts: list[tuple[str, dict]] = []

    def handle_starttag(self, tag, attrs):
        if tag == "br":
            self.parts.append(("\n", self.stack[-1]))
            return
        fmt = dict(self.stack[-1])
        attributes = dict(attrs)
        if tag in ("b", "strong"):
            fmt["bold"] = True
        elif tag in ("i", "em"):
            fmt["italic"] = True
        elif tag == "u":
            fmt["underline"] = True
        elif tag == "sub":
            fmt["vert"] = "subscript"
        elif tag in ("super", "sup"):
            fmt["vert"] = "superscript"
        elif tag == "font":
            if attributes.get("color"):
                fmt["color"] = _hex(attributes["color"]) or fmt.get("color")
            if attributes.get("size"):
                fmt["size"] = float(attributes["size"])
            face = attributes.get("face") or attributes.get("name") or ""
            if "bold" in face.lower():
                fmt["bold"] = True
        self.stack.append(fmt)

    def handle_startendtag(self, tag, attrs):
        if tag == "br":
            self.parts.append(("\n", self.stack[-1]))

    def handle_endtag(self, tag):
        if tag != "br" and len(self.stack) > 1:
            self.stack.pop()

    def handle_data(self, data):
        if data:
            self.parts.append((data.replace("\n", " "), self.stack[-1]))


def _run_properties(parent, fmt: dict):
    rpr = _el("w:rPr", parent)
    _el("w:rFonts", rpr, ascii=_FONT, hAnsi=_FONT, cs=_FONT, eastAsia=_FONT)
    if fmt.get("bold"):
        _el("w:b", rpr)
        _el("w:bCs", rpr)
    if fmt.get("italic"):
        _el("w:i", rpr)
        _el("w:iCs", rpr)
    if fmt.get("color"):
        _el("w:color", rpr, val=fmt["color"])
    size = int(round(fmt.get("size", funcef.BODY_SIZE) * 2))
    _el("w:sz", rpr, val=size)
    _el("w:szCs", rpr, val=size)
    if fmt.get("underline"):
        _el("w:u", rpr, val="single")
    if fmt.get("vert"):
        _el("w:vertAlign", rpr, val=fmt["vert"])
    return rpr


def _append_runs(paragraph, parts):
    for text, fmt in parts:
        run = _el("w:r", paragraph)
        _run_properties(run, fmt)
        for idx, piece in enumerate(text.split("\n")):
            if idx:
                _el("w:br", run)
            if piece:
                t = _el("w:t", run)
                t.text = piece
                t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def _style_format(style) -> dict:
    font = (style.fontName or "").lower()
    return {
        "bold": "bold" in font,
        "italic": "italic" in font or "oblique" in font,
        "color": _hex(style.textColor),
        "size": float(style.fontSize),
    }


# ---------------------------------------------------------------------------
# Conversor do story
# ---------------------------------------------------------------------------

def _is_decorative(cell) -> bool:
    if isinstance(cell, Drawing):
        return True
    if isinstance(cell, (list, tuple)):
        return bool(cell) and all(_is_decorative(x) for x in cell)
    return False


def _is_blank(cell) -> bool:
    return cell is None or (isinstance(cell, str) and not cell.strip())


class _StoryWriter:
    def __init__(self, pdf_width: float):
        self.scale = _CONTENT_WIDTH_TW / _tw(pdf_width)
        self.pdf_width = pdf_width
        self._canvas = Canvas(BytesIO(), pagesize=A4)

    # Elementos de bloco -------------------------------------------------
    def convert(self, flowables) -> list:
        out: list = []
        for item in flowables:
            self._flowable(item, out)
        return out

    def _flowable(self, item, out, keep_next=False):
        if isinstance(item, Paragraph):
            out.append(self.paragraph(item, keep_next=keep_next))
        elif isinstance(item, Table):
            table = self.table(item)
            if table is not None:
                keep = keep_next or bool(getattr(item, "keepWithNext", False))
                rows = table.findall(qn("w:tr"))
                if len(rows) <= _KEEP_TABLE_ROWS:
                    for tr in rows[:-1]:
                        _keep_paragraphs(tr)
                if keep:
                    _keep_paragraphs(rows[-1])
                out.append(table)
                spacer = self.spacer(1.5)
                if keep:
                    _keep_paragraphs(spacer)
                out.append(spacer)
        elif isinstance(item, KeepTogether):
            content = list(item._content)
            for idx, child in enumerate(content):
                self._flowable(child, out, keep_next=idx < len(content) - 1)
        elif isinstance(item, (list, tuple)):
            for child in item:
                self._flowable(child, out)
        elif isinstance(item, Spacer) and not isinstance(item, (CondPageBreak, PageBreak)):
            out.append(self.spacer(item.height))
        elif isinstance(item, HRFlowable):
            out.append(self.rule(item))
        elif isinstance(item, v13._KPIGrid):
            out.append(self.kpi_table(item.items))
            out.append(self.spacer(1.5))
        # Quebras condicionais, ícones e imagens decorativas não vão para o Word.

    def paragraph(self, item: Paragraph, keep_next=False, alignment=None):
        style = item.style
        p = _el("w:p")
        ppr = _el("w:pPr", p)
        if keep_next or getattr(style, "keepWithNext", 0):
            _el("w:keepNext", ppr)
        _el(
            "w:spacing", ppr,
            before=_tw(style.spaceBefore or 0), after=_tw(style.spaceAfter or 0),
            line=_tw(max(style.leading, style.fontSize * 1.15)), lineRule="atLeast",
        )
        if style.leftIndent or style.firstLineIndent or style.rightIndent:
            _el("w:ind", ppr, left=_tw(style.leftIndent or 0), right=_tw(style.rightIndent or 0),
                firstLine=_tw(max(style.firstLineIndent or 0, 0)))
        _el("w:jc", ppr, val=alignment or _ALIGN.get(style.alignment, "left"))
        mark = _el("w:rPr", ppr)
        size = int(round(style.fontSize * 2))
        _el("w:sz", mark, val=size)
        _el("w:szCs", mark, val=size)
        parser = _Markup(_style_format(style))
        parser.feed(item.text or "")
        parser.close()
        _append_runs(p, parser.parts)
        return p

    def text_paragraph(self, text: str, fmt: dict, alignment="left"):
        p = _el("w:p")
        ppr = _el("w:pPr", p)
        _el("w:spacing", ppr, before=0, after=0)
        _el("w:jc", ppr, val=alignment)
        _append_runs(p, [(text, fmt)])
        return p

    def spacer(self, height_pt: float):
        p = _el("w:p")
        ppr = _el("w:pPr", p)
        _el("w:spacing", ppr, before=0, after=0, line=max(_tw(height_pt), 20), lineRule="exact")
        mark = _el("w:rPr", ppr)
        _el("w:sz", mark, val=2)
        _el("w:szCs", mark, val=2)
        return p

    def rule(self, item: HRFlowable):
        p = _el("w:p")
        ppr = _el("w:pPr", p)
        borders = _el("w:pBdr", ppr)
        _el("w:bottom", borders, val="single", sz=max(2, int(round(item.lineWidth * 8))),
            space=1, color=_hex(item.color) or "C9D7E6")
        _el("w:spacing", ppr, before=_tw(item.spaceBefore or 0), after=_tw(item.spaceAfter or 0),
            line=20, lineRule="exact")
        return p

    # Tabelas -----------------------------------------------------------
    def table(self, table: Table):
        table.wrapOn(self._canvas, self.pdf_width, 100000)
        data = table._cellvalues
        rows = len(data)
        cols = len(data[0]) if rows else 0
        if not rows or not cols:
            return None

        def norm(value, size):
            return value + size if value < 0 else value

        def cells(start, stop):
            (sc, sr), (ec, er) = start, stop
            sc, ec = norm(sc, cols), norm(ec, cols)
            sr, er = norm(sr, rows), norm(er, rows)
            return sc, sr, min(ec, cols - 1), min(er, rows - 1)

        decorative_cols = {
            c for c in range(cols)
            if any(_is_decorative(data[r][c]) for r in range(rows))
            and all(_is_decorative(data[r][c]) or _is_blank(data[r][c]) for r in range(rows))
        }
        decorative_rows = {
            r for r in range(rows)
            if any(_is_decorative(data[r][c]) for c in range(cols))
            and all(_is_decorative(data[r][c]) or _is_blank(data[r][c]) for c in range(cols))
        }
        keep_cols = [c for c in range(cols) if c not in decorative_cols]
        keep_rows = [r for r in range(rows) if r not in decorative_rows]
        if not keep_cols or not keep_rows:
            return None

        widths = list(table._colWidths)
        for c in decorative_cols:
            # A largura do ícone passa para a coluna seguinte mantida (ou a anterior).
            target = next((k for k in keep_cols if k > c), keep_cols[-1])
            widths[target] += widths[c]
        widths_tw = [max(1, int(round(_tw(widths[c]) * self.scale))) for c in keep_cols]

        # Fundos e bordas por célula original.
        backgrounds = {}
        for cmd in table._bkgrndcmds:
            op, start, stop, arg = cmd[0], cmd[1], cmd[2], cmd[3]
            sc, sr, ec, er = cells(start, stop)
            for r in range(sr, er + 1):
                for c in range(sc, ec + 1):
                    if op == "BACKGROUND":
                        color = arg
                    elif op == "ROWBACKGROUNDS" and arg:
                        color = arg[(r - sr) % len(arg)]
                    elif op == "COLBACKGROUNDS" and arg:
                        color = arg[(c - sc) % len(arg)]
                    else:
                        continue
                    if color is not None:
                        backgrounds[(r, c)] = _hex(color)
        borders: dict = {}

        def put(r, c, side, weight, color):
            if weight and weight > 0:
                borders.setdefault((r, c), {})[side] = (weight, _hex(color) or "000000")
            else:
                borders.setdefault((r, c), {}).pop(side, None)

        for cmd in table._linecmds:
            op, start, stop, weight, color = cmd[0], cmd[1], cmd[2], cmd[3], cmd[4]
            sc, sr, ec, er = cells(start, stop)
            for r in range(sr, er + 1):
                for c in range(sc, ec + 1):
                    if op == "GRID":
                        for side in ("top", "left", "bottom", "right"):
                            put(r, c, side, weight, color)
                    elif op in ("BOX", "OUTLINE"):
                        if r == sr: put(r, c, "top", weight, color)
                        if r == er: put(r, c, "bottom", weight, color)
                        if c == sc: put(r, c, "left", weight, color)
                        if c == ec: put(r, c, "right", weight, color)
                    elif op == "INNERGRID":
                        if r > sr: put(r, c, "top", weight, color)
                        if r < er: put(r, c, "bottom", weight, color)
                        if c > sc: put(r, c, "left", weight, color)
                        if c < ec: put(r, c, "right", weight, color)
                    elif op == "LINEABOVE":
                        put(r, c, "top", weight, color)
                    elif op == "LINEBELOW":
                        put(r, c, "bottom", weight, color)
                    elif op == "LINEBEFORE":
                        put(r, c, "left", weight, color)
                    elif op == "LINEAFTER":
                        put(r, c, "right", weight, color)

        # Mesclagens.
        span_of = {}
        covered = set()
        for cmd in table._spanCmds:
            sc, sr, ec, er = cells(cmd[1], cmd[2])
            span_of[(sr, sc)] = (ec, er)
            for r in range(sr, er + 1):
                for c in range(sc, ec + 1):
                    if (r, c) != (sr, sc):
                        covered.add((r, c))
        vmerge_continue = {}
        for (sr, sc), (ec, er) in span_of.items():
            for r in range(sr + 1, er + 1):
                vmerge_continue[(r, sc)] = (sc, ec)

        tbl = _el("w:tbl")
        tblpr = _el("w:tblPr", tbl)
        _el("w:tblW", tblpr, w=sum(widths_tw), type="dxa")
        align = {"LEFT": "left", "RIGHT": "right"}.get(str(table.hAlign).upper(), "center")
        _el("w:jc", tblpr, val=align)
        _el("w:tblLayout", tblpr, type="fixed")
        margins = _el("w:tblCellMar", tblpr)
        for side in ("top", "left", "bottom", "right"):
            _el(f"w:{side}", margins, w=0, type="dxa")
        grid = _el("w:tblGrid", tbl)
        for width in widths_tw:
            _el("w:gridCol", grid, w=width)

        new_col = {c: i for i, c in enumerate(keep_cols)}
        heights = table._argH or [None] * rows
        min_heights = getattr(table, "_minRowHeights", None) or [None] * rows
        header_rows = table.repeatRows if isinstance(table.repeatRows, int) else 0

        for r in keep_rows:
            tr = _el("w:tr", tbl)
            trpr = _el("w:trPr", tr)
            _el("w:cantSplit", trpr)
            fixed = heights[r] if r < len(heights) else None
            minimum = min_heights[r] if r < len(min_heights) else None
            height = max(x for x in (fixed or 0, minimum or 0))
            if height:
                _el("w:trHeight", trpr, val=_tw(height), hRule="atLeast")
            if r < header_rows:
                _el("w:tblHeader", trpr)
            c = 0
            while c < cols:
                if c in decorative_cols or ((r, c) in covered and (r, c) not in vmerge_continue):
                    c += 1
                    continue
                if (r, c) in vmerge_continue:
                    sc, ec = vmerge_continue[(r, c)]
                    span_cols = [k for k in range(sc, ec + 1) if k in new_col]
                    tc = self._cell(tr, sum(widths_tw[new_col[k]] for k in span_cols), len(span_cols))
                    _el("w:vMerge", tc.find(qn("w:tcPr")))
                    self._finish_cell(tc, table, r, sc, ec, borders, backgrounds, continuation=True)
                    c = ec + 1
                    continue
                ec, er = span_of.get((r, c), (c, r))
                span_cols = [k for k in range(c, ec + 1) if k in new_col]
                if not span_cols:
                    c = ec + 1
                    continue
                tc = self._cell(tr, sum(widths_tw[new_col[k]] for k in span_cols), len(span_cols))
                if er > r:
                    _el("w:vMerge", tc.find(qn("w:tcPr")), val="restart")
                self._finish_cell(tc, table, r, c, ec, borders, backgrounds)
                self._cell_content(tc, data[r][c], table._cellStyles[r][c])
                c = ec + 1
        return tbl

    def _cell(self, tr, width_tw, span):
        tc = _el("w:tc", tr)
        tcpr = _el("w:tcPr", tc)
        _el("w:tcW", tcpr, w=width_tw, type="dxa")
        if span > 1:
            _el("w:gridSpan", tcpr, val=span)
        return tc

    def _finish_cell(self, tc, table, r, c, ec, borders, backgrounds, continuation=False):
        tcpr = tc.find(qn("w:tcPr"))
        first = borders.get((r, c), {})
        last = borders.get((r, ec), {})
        sides = {
            "top": first.get("top"), "left": first.get("left"),
            "bottom": first.get("bottom"), "right": last.get("right"),
        }
        tcb = _el("w:tcBorders", tcpr)
        for side in ("top", "left", "bottom", "right"):
            value = sides[side]
            if value:
                weight, color = value
                _el(f"w:{side}", tcb, val="single", sz=max(2, int(round(weight * 8))), space=0, color=color)
            else:
                _el(f"w:{side}", tcb, val="nil")
        fill = backgrounds.get((r, c))
        if fill:
            _el("w:shd", tcpr, val="clear", color="auto", fill=fill)
        style = table._cellStyles[r][c]
        mar = _el("w:tcMar", tcpr)
        _el("w:top", mar, w=_tw(style.topPadding), type="dxa")
        _el("w:left", mar, w=_tw(style.leftPadding), type="dxa")
        _el("w:bottom", mar, w=_tw(style.bottomPadding), type="dxa")
        _el("w:right", mar, w=_tw(style.rightPadding), type="dxa")
        _el("w:vAlign", tcpr, val=_VALIGN.get(str(style.valign).upper(), "top"))
        if continuation:
            tc.append(self.text_paragraph("", {"size": funcef.TABLE_SIZE}))

    def _cell_content(self, tc, value, cell_style):
        alignment = _CELL_ALIGN.get(str(cell_style.alignment).upper())
        items = value if isinstance(value, (list, tuple)) else [value]
        produced = []
        for item in items:
            if isinstance(item, Paragraph):
                produced.append(self.paragraph(item))
            elif isinstance(item, str):
                fmt = {
                    "bold": "bold" in (cell_style.fontname or "").lower(),
                    "color": _hex(cell_style.textColor),
                    "size": float(cell_style.fontsize),
                }
                produced.append(self.text_paragraph(item, fmt, alignment or "left"))
            elif item is None or isinstance(item, (Drawing, Image, CondPageBreak)):
                continue
            else:
                self._flowable(item, produced)
        if not produced or produced[-1].tag != qn("w:p"):
            produced.append(self.text_paragraph("", {"size": funcef.TABLE_SIZE}))
        for element in produced:
            tc.append(element)

    def kpi_table(self, items):
        n = len(items)
        width = _CONTENT_WIDTH_TW // n
        tbl = _el("w:tbl")
        tblpr = _el("w:tblPr", tbl)
        _el("w:tblW", tblpr, w=width * n, type="dxa")
        _el("w:jc", tblpr, val="center")
        _el("w:tblLayout", tblpr, type="fixed")
        grid = _el("w:tblGrid", tbl)
        for _ in items:
            _el("w:gridCol", grid, w=width)
        tr = _el("w:tr", tbl)
        trpr = _el("w:trPr", tr)
        _el("w:cantSplit", trpr)
        for label, value in items:
            tc = self._cell(tr, width, 1)
            tcpr = tc.find(qn("w:tcPr"))
            tcb = _el("w:tcBorders", tcpr)
            for side in ("top", "left", "bottom", "right"):
                _el(f"w:{side}", tcb, val="single", sz=4, space=0, color="C9D8E6")
            mar = _el("w:tcMar", tcpr)
            for side, w in (("top", 90), ("left", 60), ("bottom", 90), ("right", 60)):
                _el(f"w:{side}", mar, w=w, type="dxa")
            _el("w:vAlign", tcpr, val="center")
            tc.append(self.text_paragraph(label, {"size": 7.0, "color": "243748"}, "center"))
            tc.append(self.text_paragraph(value, {"size": 9.0, "bold": True, "color": "073D6D"}, "center"))
        return tbl


# ---------------------------------------------------------------------------
# Modelo corporativo
# ---------------------------------------------------------------------------

def _set_cell_value(tc, text: str) -> None:
    paragraph = tc.find(qn("w:p"))
    if paragraph is None:
        paragraph = _el("w:p", tc)
    for child in list(paragraph):
        if child.tag != qn("w:pPr"):
            paragraph.remove(child)
    _append_runs(paragraph, [(text, {"color": _VALUE_COLOR, "size": 10})])
    tcpr = tc.find(qn("w:tcPr"))
    if tcpr is not None and tcpr.find(qn("w:vAlign")) is None:
        _el("w:vAlign", tcpr, val="center")


def _fill_opening(body, header: ManifestationHeader) -> None:
    fields = {
        "Processo nº": header.process_number,
        "Comarca": header.district,
        "Vara": header.court,
        "UF": header.state,
        "Autor(es)": header.authors,
        "Advogado responsável": header.lawyer,
        "Mutuário(s)": header.borrowers,
        "Matrícula(s)": header.registrations,
        "Operação com participante": header.operation,
        "Contrato(s)": header.contracts,
        "Modalidade": header.modality,
        "Área de destino": header.destination_area,
        "Assunto": header.subject,
        "Referência": header.reference,
    }
    for tc in list(body.iter(qn("w:tc"))):
        if tc.find(".//" + qn("w:tbl")) is not None:
            continue
        label = _text_of(tc)
        if label not in fields:
            continue
        value_tc = tc.getnext()
        while value_tc is not None and value_tc.tag != qn("w:tc"):
            value_tc = value_tc.getnext()
        value = (fields[label] or "").strip()
        if value_tc is not None and value:
            _set_cell_value(value_tc, value)

    for box in body.iter(qn("w:txbxContent")):
        paragraphs = box.findall(qn("w:p"))
        if not any(_text_of(x) == "Data:" for x in paragraphs):
            continue
        for paragraph in paragraphs:
            if _text_of(paragraph) == "Data:":
                break
            box.remove(paragraph)

    issue = (header.issue_date or date.today()).strftime("%d/%m/%Y")
    for t in body.iter(qn("w:t")):
        if t.text and _TEMPLATE_DATE in t.text:
            t.text = t.text.replace(_TEMPLATE_DATE, issue)


def _responsible_table(body):
    for tbl in body.iter(qn("w:tbl")):
        if _text_of(tbl).startswith("Responsável pela informação"):
            return tbl
    raise ValueError("Modelo sem o quadro 'Responsável pela informação'.")


def _fill_responsible(tbl, name: str) -> None:
    for tr in list(tbl.iter(qn("w:tr"))):
        if _text_of(tr).startswith("Conferido por"):
            tr.getparent().remove(tr)
    paragraph = next(tbl.iter(qn("w:p")))
    name = (name or "").strip()
    if name:
        _append_runs(paragraph, [(name, {"bold": True, "color": "000000", "size": 10})])


def _add_page_numbers(document) -> None:
    footer = document.sections[0].footer
    paragraph = footer.paragraphs[0]._p
    ppr = paragraph.find(qn("w:pPr"))
    if ppr is None:
        ppr = _el("w:pPr")
        paragraph.insert(0, ppr)
    tabs = _el("w:tabs")
    _el("w:tab", tabs, val="right", pos=_CONTENT_WIDTH_TW)
    spacing = ppr.find(qn("w:spacing"))
    if spacing is not None:
        spacing.addprevious(tabs)
    else:
        ppr.append(tabs)
    jc = ppr.find(qn("w:jc"))
    if jc is not None:
        jc.set(qn("w:val"), "left")
    fmt = {"color": _FOOTER_COLOR, "size": 6}
    first_run = ppr.getnext()
    pieces = []

    def run(text=None, tab=False):
        r = _el("w:r")
        _run_properties(r, fmt)
        if tab:
            _el("w:tab", r)
        if text is not None:
            t = _el("w:t", r)
            t.text = text
            t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        return r

    def field(instr):
        begin = run()
        _el("w:fldChar", begin, fldCharType="begin")
        code = run()
        instr_text = _el("w:instrText", code)
        instr_text.text = f" {instr} "
        instr_text.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        separate = run()
        _el("w:fldChar", separate, fldCharType="separate")
        end = run()
        _el("w:fldChar", end, fldCharType="end")
        return [begin, code, separate, run("1"), end]

    pieces = [run("Página "), *field("PAGE"), run(" de "), *field("NUMPAGES"), run(tab=True)]
    for piece in pieces:
        if first_run is not None:
            first_run.addprevious(piece)
        else:
            paragraph.append(piece)


def build_opinion_docx(
    *,
    settings,
    result,
    modality_name: str,
    identity,
    selected_installment_numbers,
    overrides=None,
) -> bytes:
    header = identity.header or ManifestationHeader()
    if not header.contracts.strip() or not header.modality.strip():
        from dataclasses import replace
        header = replace(
            header,
            contracts=header.contracts.strip() or identity.contract_number.strip(),
            modality=header.modality.strip() or modality_name.strip(),
        )

    body_story = v13.opinion_body_story(
        settings=settings,
        result=result,
        modality_name=modality_name,
        identity=identity,
        selected_installment_numbers=selected_installment_numbers,
        overrides=overrides,
    )
    pdf_width = v13.CONTENT_W if settings.modality_code in {"MOD_001", "MOD_002"} else base.CONTENT_WIDTH
    elements = _StoryWriter(pdf_width).convert(body_story)

    document = Document(str(TEMPLATE_PATH))
    body = document.element.body
    _fill_opening(body, header)

    responsible = _responsible_table(body)
    _fill_responsible(responsible, header.responsible)

    # Remove os parágrafos vazios entre o quadro 03 e o responsável e insere o conteúdo técnico.
    previous = responsible.getprevious()
    while previous is not None and previous.tag == qn("w:p") and not _text_of(previous):
        to_remove = previous
        previous = previous.getprevious()
        body.remove(to_remove)
    writer = _StoryWriter(pdf_width)
    for element in [writer.spacer(14)] + elements + [writer.spacer(16)]:
        responsible.addprevious(element)

    _add_page_numbers(document)
    document.core_properties.title = f"Parecer técnico - {modality_name}"
    document.core_properties.author = "FUNCEF - Motor de Cálculos"
    document.core_properties.subject = "Evolução contratual"

    output = BytesIO()
    document.save(output)
    return output.getvalue()
