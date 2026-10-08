from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal, ROUND_HALF_UP, localcontext
from io import BytesIO
import calendar
from typing import Iterable

from .indices import IndexSeries
from .monetary_update import excel_yearfrac


MONEY = Decimal('0.01')
FACTOR_Q = Decimal('0.000001')
SEM_CORRECAO = 'SEM_CORRECAO'


def qmoney(v: Decimal) -> Decimal:
    return v.quantize(MONEY, rounding=ROUND_HALF_UP)


def qfactor(v: Decimal) -> Decimal:
    return v.quantize(FACTOR_Q, rounding=ROUND_HALF_UP)


def add_months(d: date, months: int, day: int | None = None) -> date:
    y = d.year + (d.month - 1 + months) // 12
    m = (d.month - 1 + months) % 12 + 1
    target_day = day if day is not None else d.day
    target_day = min(target_day, calendar.monthrange(y, m)[1])
    return date(y, m, target_day)


def competence_key(d: date, lag_months: int = 0) -> str:
    ref = add_months(date(d.year, d.month, 1), -lag_months, 1)
    return f'{ref.year:04d}-{ref.month:02d}'


def _month_end(d: date) -> date:
    return date(d.year, d.month, calendar.monthrange(d.year, d.month)[1])


def _month_start(d: date) -> date:
    return date(d.year, d.month, 1)


@dataclass(frozen=True)
class MAJSSettings:
    original_amount: Decimal
    credit_date: date
    monthly_interest_rate: Decimal
    term_months: int
    base_date: date
    correction_index_code: str = 'INPC'
    correction_lag_months: int = 2
    due_day: int = 20

    def __post_init__(self):
        if self.original_amount <= 0:
            raise ValueError('O valor original deve ser positivo.')
        if self.term_months <= 0:
            raise ValueError('O prazo deve ser positivo.')
        if self.base_date <= self.credit_date:
            raise ValueError('A data-base deve ser posterior à data do crédito.')
        if not 1 <= self.due_day <= 28:
            raise ValueError('O dia de vencimento MAJS deve ficar entre 1 e 28.')
        if self.correction_lag_months < 0:
            raise ValueError('A defasagem do índice não pode ser negativa.')


@dataclass(frozen=True)
class MAJSPayment:
    payment_date: date
    amount: Decimal
    installment_number: int | None = None
    competence: str | None = None
    note: str = ''


@dataclass(frozen=True)
class MAJSExtraAmortization:
    event_date: date
    amount: Decimal
    note: str = ''


@dataclass(frozen=True)
class MAJSSuspension:
    start_date: date
    end_date: date
    note: str = ''

    def __post_init__(self):
        if self.end_date < self.start_date:
            raise ValueError('Na suspensão, a data final não pode ser anterior à data inicial.')


@dataclass(frozen=True)
class MAJSEarlySettlement:
    settlement_date: date
    novation: bool
    amount_paid: Decimal = Decimal('0')
    note: str = ''

    def __post_init__(self):
        if self.amount_paid < 0:
            raise ValueError('O valor pago na quitação não pode ser negativo.')
        if self.novation and self.amount_paid != 0:
            # Na novação não há desembolso do participante a conciliar como pagamento.
            object.__setattr__(self, 'amount_paid', Decimal('0'))


@dataclass(frozen=True)
class MAJSRow:
    installment_number: int | None
    due_date: date
    opening_balance: Decimal
    factor_position: int
    financial_factor: Decimal
    remaining_factor_sum: Decimal
    installment_due: Decimal
    interest_component: Decimal
    amortization_component: Decimal
    correction_reference: str
    correction_rate: Decimal
    extra_amortization: Decimal
    closing_balance: Decimal
    event_type: str = 'PRESTACAO'
    event_note: str = ''
    excess_extra_amortization: Decimal = Decimal('0')
    settlement_payment: Decimal = Decimal('0')


@dataclass(frozen=True)
class MAJSDifferenceRow:
    installment_number: int | None
    due_date: date
    amount_due: Decimal
    amount_paid: Decimal
    difference_paid_minus_due: Decimal
    correction_factor: Decimal = Decimal('1')
    corrected_difference: Decimal = Decimal('0')
    accumulated_interest_rate: Decimal = Decimal('0')
    interest_amount: Decimal = Decimal('0')
    updated_difference: Decimal = Decimal('0')
    interest_start_date: date | None = None
    updated_through_date: date | None = None
    origin: str = 'Prestação'
    note: str = ''


@dataclass(frozen=True)
class _ExcessAmortization:
    event_date: date
    amount: Decimal
    note: str = ''


@dataclass
class MAJSResult:
    settings: MAJSSettings
    schedule: list[MAJSRow]
    differences: list[MAJSDifferenceRow]
    warnings: list[str] = field(default_factory=list)
    difference_correction_mode: str = 'INPC_IPCA_14905'
    difference_interest_mode: str = 'UM_PCT_ATE_TL'
    fixed_interest_start: date | None = None
    yearfrac_basis: int = 0
    early_settlement: MAJSEarlySettlement | None = None

    @property
    def due_total(self) -> Decimal:
        return sum((r.amount_due for r in self.differences), Decimal('0'))

    @property
    def paid_total(self) -> Decimal:
        return sum((r.amount_paid for r in self.differences), Decimal('0'))

    @property
    def difference_total(self) -> Decimal:
        return sum((r.difference_paid_minus_due for r in self.differences), Decimal('0'))

    @property
    def updated_difference_total(self) -> Decimal:
        return sum((r.updated_difference for r in self.differences), Decimal('0'))

    @property
    def early_settlement_balance(self) -> Decimal:
        rows = [r for r in self.schedule if r.event_type == 'QUITACAO_ANTECIPADA']
        return qmoney(rows[-1].installment_due) if rows else Decimal('0')

    @property
    def was_settled_early(self) -> bool:
        return bool(self.early_settlement and any(r.event_type == 'QUITACAO_ANTECIPADA' for r in self.schedule))

    @property
    def balance_to_mature(self) -> Decimal:
        if self.was_settled_early:
            return Decimal('0')
        prior = [r for r in self.schedule if r.due_date <= self.settings.base_date]
        if prior:
            return prior[-1].closing_balance
        return self.schedule[0].opening_balance if self.schedule else Decimal('0')

    @property
    def last_evolution_date(self) -> date | None:
        return self.schedule[-1].due_date if self.schedule else None


@dataclass(frozen=True)
class _PartialFactor:
    factor: Decimal
    effective_end: date
    missing_code: str | None = None
    missing_key: str | None = None


@dataclass(frozen=True)
class _PartialInterest:
    rate: Decimal
    effective_end: date
    missing_code: str | None = None
    missing_key: str | None = None


def _first_due_date(credit_date: date, due_day: int) -> date:
    return add_months(date(credit_date.year, credit_date.month, 1), 1, due_day)


def _factor(monthly_rate: Decimal, position: int) -> Decimal:
    with localcontext() as ctx:
        ctx.prec = 40
        return qfactor(Decimal('1') / (Decimal('1') + monthly_rate * Decimal(position)))


def _factor_sum(monthly_rate: Decimal, count: int) -> Decimal:
    return qfactor(sum((_factor(monthly_rate, n) for n in range(1, count + 1)), Decimal('0')))


def _is_suspended(due: date, suspensions: Iterable[MAJSSuspension]) -> tuple[bool, str]:
    notes: list[str] = []
    active = False
    for item in suspensions:
        if item.start_date <= due <= item.end_date:
            active = True
            if item.note.strip():
                notes.append(item.note.strip())
    return active, ' | '.join(dict.fromkeys(notes))


def _schedule_correction(settings: MAJSSettings, indices: dict[str, IndexSeries], due: date) -> tuple[str, Decimal] | None:
    if settings.correction_index_code in {'', SEM_CORRECAO, None}:
        return '', Decimal('0')
    series = indices.get(settings.correction_index_code)
    if series is None:
        return None
    ref = competence_key(due, settings.correction_lag_months)
    if ref not in series.values:
        return None
    return ref, series.values[ref]


def _proportional_correction_rate(
    monthly_rate: Decimal,
    segment_start: date,
    segment_end: date,
    period_start: date,
    period_end: date,
) -> tuple[Decimal, int, int]:
    """Prorrateia o índice mensal por equivalência diária dentro do ciclo MAJS.

    O período completo é o intervalo entre o vencimento anterior e o próximo
    vencimento. Para uma amortização extraordinária no meio do ciclo, aplica-se
    somente a fração equivalente aos dias transcorridos desde o último ponto já
    corrigido. As frações sucessivas recompõem exatamente o índice mensal completo
    quando não há alteração de base.
    """
    total_days = (period_end - period_start).days
    segment_days = (segment_end - segment_start).days
    if monthly_rate == 0 or total_days <= 0 or segment_days <= 0:
        return Decimal('0'), max(segment_days, 0), max(total_days, 0)
    if monthly_rate <= Decimal('-1'):
        raise ValueError('O índice mensal deve ser superior a -100% para cálculo proporcional.')
    with localcontext() as ctx:
        ctx.prec = 40
        exponent = Decimal(segment_days) / Decimal(total_days)
        rate = (Decimal('1') + monthly_rate) ** exponent - Decimal('1')
    return rate, segment_days, total_days


def _apply_extra_amount(balance: Decimal, amount: Decimal) -> tuple[Decimal, Decimal, Decimal]:
    applied = qmoney(min(balance, amount))
    excess = qmoney(max(Decimal('0'), amount - applied))
    return qmoney(balance - applied), applied, excess


def _calculate_majs_schedule_detailed(
    settings: MAJSSettings,
    indices: dict[str, IndexSeries],
    extra_amortizations: Iterable[MAJSExtraAmortization] = (),
    suspensions: Iterable[MAJSSuspension] = (),
    early_settlement: MAJSEarlySettlement | None = None,
) -> tuple[list[MAJSRow], list[str], list[_ExcessAmortization]]:
    """Reconstrói o fluxo MAJS preservando a ordem econômica dos eventos.

    Regra para amortização extraordinária:
    - entre vencimentos, aplica-se primeiro somente a correção proporcional aos dias
      transcorridos até a data da amortização, por equivalência diária;
    - a amortização extraordinária é aplicada sobre esse saldo parcialmente corrigido;
    - no vencimento seguinte aplica-se apenas a parcela remanescente da correção do ciclo;
    - as parcelas de correção recompõem exatamente o índice mensal completo quando não há
      alteração da base;
    - a amortização reinicia a posição do fator financeiro.
    """
    if settings.correction_index_code not in {'', SEM_CORRECAO, None} and settings.correction_index_code not in indices:
        raise KeyError(f'Índice {settings.correction_index_code} não disponível para o MAJS.')

    if early_settlement is not None:
        if early_settlement.settlement_date <= settings.credit_date:
            raise ValueError('A data da quitação antecipada deve ser posterior à data do crédito.')
        if early_settlement.settlement_date > settings.base_date:
            raise ValueError('A data da quitação antecipada não pode ser posterior à data-base.')
        if not early_settlement.novation and early_settlement.amount_paid < 0:
            raise ValueError('Informe um valor pago válido para a quitação antecipada.')

    extras = sorted(list(extra_amortizations), key=lambda x: (x.event_date, x.amount))
    suspension_rows = sorted(list(suspensions), key=lambda x: (x.start_date, x.end_date))
    if any(x.amount <= 0 for x in extras):
        raise ValueError('Amortizações extraordinárias devem ser positivas.')

    first_due = _first_due_date(settings.credit_date, settings.due_day)
    grace_days = (first_due - settings.credit_date).days
    opening = qmoney(
        settings.original_amount
        + settings.original_amount * (settings.monthly_interest_rate / Decimal('30')) * Decimal(grace_days)
    )

    rows: list[MAJSRow] = []
    warnings: list[str] = []
    excess_events: list[_ExcessAmortization] = []
    processed_extra: set[int] = set()
    factor_position = 1
    balance = opening
    previous_due = settings.credit_date
    installment_no = 1
    due = first_due
    suspended_count = 0

    def append_extra_event(
        *,
        extra: MAJSExtraAmortization,
        opening_balance: Decimal,
        ref: str,
        correction_rate_used: Decimal,
        corrected_before: Decimal,
        applied: Decimal,
        excess: Decimal,
        correction_days: int = 0,
        cycle_days: int = 0,
    ) -> None:
        note_parts = []
        if correction_rate_used != 0:
            if cycle_days > 0:
                note_parts.append(
                    f'Correção proporcional aplicada antes da amortização: {correction_days}/{cycle_days} dia(s) do ciclo.'
                )
            else:
                note_parts.append('Saldo corrigido antes da amortização extraordinária.')
        elif settings.correction_index_code in {'', SEM_CORRECAO, None}:
            note_parts.append('Modalidade sem correção monetária.')
        else:
            note_parts.append('Correção da competência já aplicada anteriormente.')
        if extra.note.strip():
            note_parts.append(extra.note.strip())
        rows.append(MAJSRow(
            installment_number=None,
            due_date=extra.event_date,
            opening_balance=opening_balance,
            factor_position=factor_position,
            financial_factor=Decimal('0'),
            remaining_factor_sum=Decimal('0'),
            installment_due=Decimal('0'),
            interest_component=Decimal('0'),
            amortization_component=Decimal('0'),
            correction_reference=ref,
            correction_rate=correction_rate_used,
            extra_amortization=applied,
            closing_balance=qmoney(corrected_before - applied),
            event_type='AMORTIZACAO_EXTRA',
            event_note=' | '.join(note_parts),
            excess_extra_amortization=excess,
        ))

    # O prazo é expresso em prestações. Competências suspensas não consomem prestação,
    # portanto o calendário pode avançar além de term_months meses.
    while (
        installment_no <= settings.term_months
        and due <= settings.base_date
        and (early_settlement is None or due < early_settlement.settlement_date)
    ):
        corr = _schedule_correction(settings, indices, due)
        if corr is None:
            ref_missing = competence_key(due, settings.correction_lag_months)
            warnings.append(
                f'Cálculo MAJS parcial: a evolução foi concluída até {previous_due.strftime("%d/%m/%Y")}. '
                f'O índice {settings.correction_index_code} necessário para {ref_missing} ainda não está disponível; '
                f'a competência de {due.strftime("%m/%Y")} não foi projetada.'
            )
            break
        ref, correction_rate = corr
        correction_cursor = previous_due

        # Amortizações entre vencimentos: aplica somente a parcela da correção
        # correspondente aos dias transcorridos até cada evento. O saldo pós-amortização
        # segue evoluindo e recebe apenas a parcela remanescente no vencimento.
        between_items = [
            (idx, extra) for idx, extra in enumerate(extras)
            if idx not in processed_extra and previous_due < extra.event_date < due
        ]
        for idx, extra in between_items:
            processed_extra.add(idx)
            event_opening = balance
            rate_used, correction_days, cycle_days = _proportional_correction_rate(
                correction_rate, correction_cursor, extra.event_date, previous_due, due
            )
            corrected_balance = qmoney(balance * (Decimal('1') + rate_used))
            after, applied, excess = _apply_extra_amount(corrected_balance, qmoney(extra.amount))
            append_extra_event(
                extra=extra,
                opening_balance=event_opening,
                ref=ref,
                correction_rate_used=rate_used,
                corrected_before=corrected_balance,
                applied=applied,
                excess=excess,
                correction_days=correction_days,
                cycle_days=cycle_days,
            )
            balance = after
            correction_cursor = extra.event_date
            if excess > 0:
                excess_events.append(_ExcessAmortization(extra.event_date, excess, extra.note))
            if applied > 0:
                factor_position = 1
            if balance <= 0:
                break

        if balance <= 0:
            break

        suspended, suspension_note = _is_suspended(due, suspension_rows)
        same_day_items = [
            (idx, extra) for idx, extra in enumerate(extras)
            if idx not in processed_extra and extra.event_date == due
        ]

        if suspended:
            suspended_count += 1
            row_opening = balance
            # Em suspensão não há prestação nem juros. O saldo é corrigido e, se houver
            # amortização extraordinária na mesma data, a correção vem primeiro.
            rate_used, _, _ = _proportional_correction_rate(
                correction_rate, correction_cursor, due, previous_due, due
            )
            balance = qmoney(balance * (Decimal('1') + rate_used))
            correction_cursor = due
            same_day_applied = Decimal('0')
            same_day_excess = Decimal('0')
            notes: list[str] = []
            if suspension_note:
                notes.append(suspension_note)
            for idx, extra in same_day_items:
                processed_extra.add(idx)
                balance, applied, excess = _apply_extra_amount(balance, qmoney(extra.amount))
                same_day_applied += applied
                same_day_excess += excess
                if extra.note.strip():
                    notes.append(extra.note.strip())
                if excess > 0:
                    excess_events.append(_ExcessAmortization(extra.event_date, excess, extra.note))
                if applied > 0:
                    factor_position = 1
            rows.append(MAJSRow(
                installment_number=None,
                due_date=due,
                opening_balance=row_opening,
                factor_position=factor_position,
                financial_factor=Decimal('0'),
                remaining_factor_sum=Decimal('0'),
                installment_due=Decimal('0'),
                interest_component=Decimal('0'),
                amortization_component=Decimal('0'),
                correction_reference=ref,
                correction_rate=rate_used,
                extra_amortization=qmoney(same_day_applied),
                closing_balance=balance,
                event_type='SUSPENSAO',
                event_note=' | '.join(dict.fromkeys(notes)),
                excess_extra_amortization=qmoney(same_day_excess),
            ))
            previous_due = due
            due = add_months(due, 1, settings.due_day)
            if balance <= 0:
                break
            continue

        # Prestação ordinária MAJS. Mantemos a metodologia validada: cálculo da
        # prestação/juros/amortização contratual sobre o saldo de abertura. A correção
        # da competência ocorre após a amortização contratual. Se houver amortização
        # extraordinária na mesma data, ela ocorre somente depois dessa correção.
        row_opening = balance
        remaining_count = settings.term_months - installment_no + 1
        pos = factor_position
        factors = [_factor(settings.monthly_interest_rate, n) for n in range(pos, pos + remaining_count)]
        factor_sum = qfactor(sum(factors, Decimal('0')))
        current_factor = factors[0]
        installment = qmoney(balance / factor_sum) if factor_sum else Decimal('0')
        interest = qmoney(installment * (Decimal('1') - current_factor))
        amortization = qmoney(min(balance, installment - interest))
        balance = qmoney(max(Decimal('0'), balance - amortization))

        rate_used, _, _ = _proportional_correction_rate(
            correction_rate, correction_cursor, due, previous_due, due
        )
        balance = qmoney(balance * (Decimal('1') + rate_used))
        correction_cursor = due

        same_day_applied = Decimal('0')
        same_day_excess = Decimal('0')
        same_day_notes: list[str] = []
        for idx, extra in same_day_items:
            processed_extra.add(idx)
            balance, applied, excess = _apply_extra_amount(balance, qmoney(extra.amount))
            same_day_applied += applied
            same_day_excess += excess
            if extra.note.strip():
                same_day_notes.append(extra.note.strip())
            if excess > 0:
                excess_events.append(_ExcessAmortization(extra.event_date, excess, extra.note))

        rows.append(MAJSRow(
            installment_number=installment_no,
            due_date=due,
            opening_balance=row_opening,
            factor_position=pos,
            financial_factor=current_factor,
            remaining_factor_sum=factor_sum,
            installment_due=installment,
            interest_component=interest,
            amortization_component=amortization,
            correction_reference=ref,
            correction_rate=rate_used,
            extra_amortization=qmoney(same_day_applied),
            closing_balance=balance,
            event_type='PRESTACAO',
            event_note=' | '.join(dict.fromkeys(same_day_notes)),
            excess_extra_amortization=qmoney(same_day_excess),
        ))

        previous_due = due
        if same_day_applied > 0:
            factor_position = 1
        else:
            factor_position += 1
        installment_no += 1
        due = add_months(due, 1, settings.due_day)
        if balance <= 0:
            break

    # Quitação antecipada encerra o fluxo contratual na data informada. Prestação com
    # vencimento exatamente nessa data não é gerada separadamente: o evento de quitação
    # substitui o restante do fluxo e registra o saldo teórico devido naquela posição.
    if early_settlement is not None and balance > 0:
        settlement_date = early_settlement.settlement_date
        next_due = due

        # Se a quitação ocorrer depois da última prestação contratual já processada, não
        # existe saldo a liquidar. Caso contrário, calculamos a posição dentro do ciclo.
        if installment_no <= settings.term_months and settlement_date >= previous_due:
            corr = _schedule_correction(settings, indices, next_due)
            if corr is None:
                ref_missing = competence_key(next_due, settings.correction_lag_months)
                raise ValueError(
                    f'Não foi possível apurar o saldo de quitação em {settlement_date.strftime("%d/%m/%Y")}: '
                    f'o índice {settings.correction_index_code} necessário para {ref_missing} não está disponível.'
                )
            ref, correction_rate = corr
            correction_cursor = previous_due

            # Amortizações extraordinárias anteriores à quitação continuam produzindo
            # efeito econômico. Eventos na mesma data ou posteriores são ignorados para
            # evitar duplicidade com o valor específico da quitação.
            before_settlement = [
                (idx, extra) for idx, extra in enumerate(extras)
                if idx not in processed_extra and previous_due < extra.event_date < settlement_date
            ]
            for idx, extra in before_settlement:
                processed_extra.add(idx)
                event_opening = balance
                rate_used, correction_days, cycle_days = _proportional_correction_rate(
                    correction_rate, correction_cursor, extra.event_date, previous_due, next_due
                )
                corrected_balance = qmoney(balance * (Decimal('1') + rate_used))
                after, applied, excess = _apply_extra_amount(corrected_balance, qmoney(extra.amount))
                append_extra_event(
                    extra=extra, opening_balance=event_opening, ref=ref,
                    correction_rate_used=rate_used, corrected_before=corrected_balance,
                    applied=applied, excess=excess, correction_days=correction_days, cycle_days=cycle_days,
                )
                balance = after
                correction_cursor = extra.event_date
                if excess > 0:
                    excess_events.append(_ExcessAmortization(extra.event_date, excess, extra.note))
                if applied > 0:
                    factor_position = 1
                if balance <= 0:
                    break

            if balance > 0:
                row_opening = balance
                rate_used, correction_days, cycle_days = _proportional_correction_rate(
                    correction_rate, correction_cursor, settlement_date, previous_due, next_due
                )
                settlement_balance = qmoney(balance * (Decimal('1') + rate_used))
                settlement_kind = 'Novação' if early_settlement.novation else 'Pagamento pelo(a) participante'
                note_parts = [
                    f'Quitação antecipada por {settlement_kind.lower()}.',
                    f'Saldo MAJS teórico apurado na data: R$ {settlement_balance:.2f}.',
                ]
                if rate_used != 0:
                    note_parts.append(
                        f'Correção proporcional até a quitação: {correction_days}/{cycle_days} dia(s) do ciclo.'
                    )
                if early_settlement.novation:
                    note_parts.append(
                        'A novação encerra o fluxo, mas não é tratada como desembolso do participante e não compõe a tabela de diferenças.'
                    )
                    settlement_payment = Decimal('0')
                else:
                    settlement_payment = qmoney(early_settlement.amount_paid)
                    note_parts.append(
                        f'Valor efetivamente pago pelo(a) participante: R$ {settlement_payment:.2f}; a diferença contra o saldo teórico compõe a tabela de diferenças.'
                    )
                if early_settlement.note.strip():
                    note_parts.append(early_settlement.note.strip())

                rows.append(MAJSRow(
                    installment_number=None,
                    due_date=settlement_date,
                    opening_balance=row_opening,
                    factor_position=factor_position,
                    financial_factor=Decimal('0'),
                    remaining_factor_sum=Decimal('0'),
                    installment_due=settlement_balance,
                    interest_component=Decimal('0'),
                    amortization_component=Decimal('0'),
                    correction_reference=ref,
                    correction_rate=rate_used,
                    extra_amortization=Decimal('0'),
                    closing_balance=Decimal('0'),
                    event_type='QUITACAO_ANTECIPADA',
                    event_note=' | '.join(note_parts),
                    excess_extra_amortization=Decimal('0'),
                    settlement_payment=settlement_payment,
                ))
                balance = Decimal('0')
                warnings.append(
                    f'O fluxo MAJS foi encerrado antecipadamente em {settlement_date.strftime("%d/%m/%Y")} por {settlement_kind.lower()}; '
                    f'não foram geradas prestações posteriores à quitação.'
                )

            ignored = [x for idx, x in enumerate(extras) if idx not in processed_extra and x.event_date >= settlement_date]
            if ignored:
                warnings.append(
                    f'{len(ignored)} amortização(ões) extraordinária(s) na data ou após a quitação foram desconsideradas. '
                    'Para quitação sem novação, informe o desembolso no campo específico de valor pago na quitação.'
                )
                processed_extra.update(idx for idx, x in enumerate(extras) if x.event_date >= settlement_date)

    # Se a dívida já foi quitada sem um evento formal de quitação antecipada, qualquer
    # amortização extraordinária posterior até a data-base é integralmente excedente.
    if balance <= 0 and early_settlement is None:
        for idx, extra in enumerate(extras):
            if idx in processed_extra or extra.event_date > settings.base_date:
                continue
            processed_extra.add(idx)
            excess_events.append(_ExcessAmortization(extra.event_date, qmoney(extra.amount), extra.note))

    if suspended_count:
        warnings.append(
            f'Foram processadas {suspended_count} competência(s) em suspensão: nelas o saldo foi apenas corrigido, '
            f'sem geração de prestação, juros ou amortização contratual. O prazo em número de prestações foi preservado.'
        )
    if excess_events:
        total_excess = qmoney(sum((x.amount for x in excess_events), Decimal('0')))
        warnings.append(
            f'Amortizações extraordinárias superaram o saldo teórico em R$ {total_excess:.2f}. '
            f'O excedente não foi descartado: foi tratado como diferença credora (pago - devido).'
        )
    return rows, warnings, excess_events

def calculate_majs_schedule(
    settings: MAJSSettings,
    indices: dict[str, IndexSeries],
    extra_amortizations: Iterable[MAJSExtraAmortization] = (),
    suspensions: Iterable[MAJSSuspension] = (),
    early_settlement: MAJSEarlySettlement | None = None,
) -> tuple[list[MAJSRow], list[str]]:
    rows, warnings, _ = _calculate_majs_schedule_detailed(settings, indices, extra_amortizations, suspensions, early_settlement)
    return rows, warnings


def _payment_competence(p: MAJSPayment) -> str:
    raw = (p.competence or '').strip()
    if raw:
        for fmt in ('%Y-%m', '%m/%Y', '%Y/%m', '%d/%m/%Y'):
            try:
                parsed = date.fromisoformat(raw + '-01') if fmt == '%Y-%m' else None
                if parsed:
                    return f'{parsed.year:04d}-{parsed.month:02d}'
            except ValueError:
                pass
            try:
                from datetime import datetime
                parsed_dt = datetime.strptime(raw, fmt)
                return f'{parsed_dt.year:04d}-{parsed_dt.month:02d}'
            except ValueError:
                pass
    return f'{p.payment_date.year:04d}-{p.payment_date.month:02d}'


def reconcile_payments(schedule: list[MAJSRow], payments: Iterable[MAJSPayment], base_date: date) -> dict[int, Decimal]:
    due_rows = [r for r in schedule if r.due_date <= base_date and r.installment_number is not None and r.event_type == 'PRESTACAO']
    matched = {int(r.installment_number): Decimal('0') for r in due_rows}
    unassigned: list[MAJSPayment] = []
    by_no = {int(r.installment_number): r for r in due_rows}

    for p in sorted(payments, key=lambda x: x.payment_date):
        if p.amount <= 0 or p.payment_date > base_date:
            continue
        if p.installment_number and p.installment_number in by_no:
            matched[p.installment_number] += p.amount
            continue
        exact = [r for r in due_rows if r.due_date == p.payment_date]
        if exact:
            matched[int(exact[0].installment_number)] += p.amount
            continue
        comp = _payment_competence(p)
        same_comp = [r for r in due_rows if f'{r.due_date.year:04d}-{r.due_date.month:02d}' == comp]
        if same_comp:
            matched[int(same_comp[0].installment_number)] += p.amount
            continue
        unassigned.append(p)

    cursor = 0
    for p in unassigned:
        if not due_rows:
            break
        while cursor < len(due_rows) - 1 and matched[int(due_rows[cursor].installment_number)] >= due_rows[cursor].installment_due:
            cursor += 1
        matched[int(due_rows[cursor].installment_number)] += p.amount
    return {k: qmoney(v) for k, v in matched.items()}


def _monthly_correction_factor_partial(
    start: date,
    end: date,
    indices: dict[str, IndexSeries],
    mode: str,
) -> _PartialFactor:
    if end <= start or mode == SEM_CORRECAO:
        return _PartialFactor(Decimal('1'), end)
    factor = Decimal('1')
    cursor = add_months(_month_start(start), 1, 1)
    final_month = _month_start(end)
    effective_end = end
    while cursor <= final_month:
        key = f'{cursor.year:04d}-{cursor.month:02d}'
        code = 'INPC' if mode == 'INPC_IPCA_14905' and cursor < date(2024, 9, 1) else ('IPCA' if mode == 'INPC_IPCA_14905' else mode)
        series = indices.get(code)
        if series is None or key not in series.values:
            previous_month = add_months(cursor, -1, 1)
            cutoff = min(end, _month_end(previous_month))
            cutoff = max(start, cutoff)
            return _PartialFactor(factor, cutoff, code, key)
        factor *= Decimal('1') + series.values[key]
        cursor = add_months(cursor, 1, 1)
    return _PartialFactor(factor, effective_end)


def _legal_rate_partial(series: IndexSeries | None, start: date, end: date) -> _PartialInterest:
    effective_start = max(start, date(2024, 8, 30))
    if end <= effective_start:
        return _PartialInterest(Decimal('0'), end)
    if series is None:
        return _PartialInterest(Decimal('0'), effective_start, 'TAXA_LEGAL', f'{effective_start.year:04d}-{effective_start.month:02d}')
    total = Decimal('0')
    cursor = effective_start
    while cursor < end:
        key = f'{cursor.year:04d}-{cursor.month:02d}'
        if key not in series.values:
            return _PartialInterest(total, cursor, 'TAXA_LEGAL', key)
        month_end_exclusive = date(cursor.year + 1, 1, 1) if cursor.month == 12 else date(cursor.year, cursor.month + 1, 1)
        segment_end = min(end, month_end_exclusive)
        days = (segment_end - cursor).days
        days_in_month = calendar.monthrange(cursor.year, cursor.month)[1]
        monthly_rate = series.values[key]
        proportional_pct = (monthly_rate * Decimal('100') * Decimal(days) / Decimal(days_in_month)).quantize(
            Decimal('0.000001'), rounding=ROUND_HALF_UP
        )
        total += proportional_pct / Decimal('100')
        cursor = segment_end
    return _PartialInterest(total, end)


def _difference_interest_rate_partial(
    start: date,
    end: date,
    indices: dict[str, IndexSeries],
    mode: str,
    yearfrac_basis: int = 0,
) -> _PartialInterest:
    if end <= start or mode == 'SEM_JUROS':
        return _PartialInterest(Decimal('0'), end)
    if mode == 'UM_PCT':
        # Juros simples: 1% a.m. apropriado pela razão dias/30.
        rate = Decimal('0.01') * (Decimal((end - start).days) / Decimal('30'))
        return _PartialInterest(rate, end)
    if mode == 'UM_PCT_YEARFRAC':
        # Mesma taxa nominal de 1% a.m. expressa como 12% a.a. e apropriada pela fração de ano
        # apurada segundo a convenção de contagem de dias selecionada.
        rate = Decimal('0.12') * excel_yearfrac(start, end, int(yearfrac_basis))
        return _PartialInterest(rate, end)
    if mode == 'TAXA_LEGAL':
        return _legal_rate_partial(indices.get('TAXA_LEGAL'), start, end)
    if mode in {'UM_PCT_ATE_TL', 'UM_PCT_YEARFRAC_ATE_TL'}:
        cut = date(2024, 8, 30)
        total = Decimal('0')
        if start < cut:
            pre_end = min(end, cut)
            if mode == 'UM_PCT_YEARFRAC_ATE_TL':
                total += Decimal('0.12') * excel_yearfrac(start, pre_end, int(yearfrac_basis))
            else:
                total += Decimal('0.01') * (Decimal((pre_end - start).days) / Decimal('30'))
        if end > cut:
            legal = _legal_rate_partial(indices.get('TAXA_LEGAL'), max(start, cut), end)
            total += legal.rate
            return _PartialInterest(total, legal.effective_end, legal.missing_code, legal.missing_key)
        return _PartialInterest(total, end)
    raise ValueError('Modo de juros das diferenças inválido.')


def _append_partial_warning(warnings: list[str], seen: set[tuple[str, str, date]], *, code: str | None, key: str | None, end: date, context: str) -> None:
    if not code or not key:
        return
    token = (code, key, end)
    if token in seen:
        return
    seen.add(token)
    warnings.append(
        f'{context} foi limitada até {end.strftime("%d/%m/%Y")} porque {code} não possui a competência {key}. '
        f'O sistema não projetou índices ausentes.'
    )


def _build_difference(
    *,
    installment_number: int | None,
    origin_date: date,
    amount_due: Decimal,
    amount_paid: Decimal,
    origin: str,
    note: str,
    settings: MAJSSettings,
    indices: dict[str, IndexSeries],
    difference_correction_mode: str,
    difference_interest_mode: str,
    fixed_interest_start: date | None,
    yearfrac_basis: int,
    warnings: list[str],
    seen_warnings: set[tuple[str, str, date]],
) -> MAJSDifferenceRow:
    diff = qmoney(amount_paid - amount_due)
    corr = _monthly_correction_factor_partial(origin_date, settings.base_date, indices, difference_correction_mode)
    _append_partial_warning(
        warnings, seen_warnings, code=corr.missing_code, key=corr.missing_key, end=corr.effective_end,
        context='A atualização monetária das diferenças',
    )
    corrected = qmoney(diff * corr.factor)

    interest_start = origin_date
    if fixed_interest_start:
        interest_start = max(interest_start, fixed_interest_start)
    interest_end = min(settings.base_date, corr.effective_end)
    interest = _difference_interest_rate_partial(
        interest_start, interest_end, indices, difference_interest_mode, yearfrac_basis
    )
    _append_partial_warning(
        warnings, seen_warnings, code=interest.missing_code, key=interest.missing_key, end=interest.effective_end,
        context='A incidência de juros sobre as diferenças',
    )
    int_amount = qmoney(corrected * interest.rate)
    updated = qmoney(corrected + int_amount)
    effective_end = min(corr.effective_end, interest.effective_end)
    return MAJSDifferenceRow(
        installment_number=installment_number,
        due_date=origin_date,
        amount_due=qmoney(amount_due),
        amount_paid=qmoney(amount_paid),
        difference_paid_minus_due=diff,
        correction_factor=corr.factor,
        corrected_difference=corrected,
        accumulated_interest_rate=interest.rate,
        interest_amount=int_amount,
        updated_difference=updated,
        interest_start_date=(None if difference_interest_mode == 'SEM_JUROS' else interest_start),
        updated_through_date=effective_end,
        origin=origin,
        note=note,
    )


def calculate_majs(
    settings: MAJSSettings,
    indices: dict[str, IndexSeries],
    payments: Iterable[MAJSPayment] = (),
    extra_amortizations: Iterable[MAJSExtraAmortization] = (),
    suspensions: Iterable[MAJSSuspension] = (),
    *,
    early_settlement: MAJSEarlySettlement | None = None,
    difference_correction_mode: str = 'INPC_IPCA_14905',
    difference_interest_mode: str = 'UM_PCT_ATE_TL',
    fixed_interest_start: date | None = None,
    yearfrac_basis: int = 0,
) -> MAJSResult:
    schedule, warnings, excess_events = _calculate_majs_schedule_detailed(
        settings, indices, extra_amortizations, suspensions, early_settlement
    )
    payments_list = list(payments)
    if early_settlement is not None:
        ordinary_payments = [p for p in payments_list if p.payment_date < early_settlement.settlement_date]
        ignored_payments = [p for p in payments_list if p.payment_date >= early_settlement.settlement_date]
        if ignored_payments:
            warnings.append(
                f'{len(ignored_payments)} pagamento(s) ordinário(s) na data ou após a quitação não foram conciliados às prestações. '
                'Na quitação sem novação, use o campo específico de valor pago na quitação.'
            )
    else:
        ordinary_payments = payments_list
    matched = reconcile_payments(schedule, ordinary_payments, settings.base_date)
    differences: list[MAJSDifferenceRow] = []
    seen_warnings: set[tuple[str, str, date]] = set()

    for row in schedule:
        if row.due_date > settings.base_date or row.installment_number is None or row.event_type != 'PRESTACAO':
            continue
        paid = matched.get(int(row.installment_number), Decimal('0'))
        differences.append(_build_difference(
            installment_number=row.installment_number,
            origin_date=row.due_date,
            amount_due=row.installment_due,
            amount_paid=paid,
            origin='Prestação',
            note='',
            settings=settings,
            indices=indices,
            difference_correction_mode=difference_correction_mode,
            difference_interest_mode=difference_interest_mode,
            fixed_interest_start=fixed_interest_start,
            yearfrac_basis=int(yearfrac_basis),
            warnings=warnings,
            seen_warnings=seen_warnings,
        ))

    for excess in excess_events:
        differences.append(_build_difference(
            installment_number=None,
            origin_date=excess.event_date,
            amount_due=Decimal('0'),
            amount_paid=excess.amount,
            origin='Excesso de amortização extraordinária',
            note=excess.note,
            settings=settings,
            indices=indices,
            difference_correction_mode=difference_correction_mode,
            difference_interest_mode=difference_interest_mode,
            fixed_interest_start=fixed_interest_start,
            yearfrac_basis=int(yearfrac_basis),
            warnings=warnings,
            seen_warnings=seen_warnings,
        ))

    settlement_row = next((r for r in schedule if r.event_type == 'QUITACAO_ANTECIPADA'), None)
    if early_settlement is not None and settlement_row is not None and not early_settlement.novation:
        differences.append(_build_difference(
            installment_number=None,
            origin_date=early_settlement.settlement_date,
            amount_due=settlement_row.installment_due,
            amount_paid=qmoney(early_settlement.amount_paid),
            origin='Quitação antecipada',
            note=early_settlement.note or 'Diferença entre o saldo MAJS teórico e o valor efetivamente pago na quitação.',
            settings=settings,
            indices=indices,
            difference_correction_mode=difference_correction_mode,
            difference_interest_mode=difference_interest_mode,
            fixed_interest_start=fixed_interest_start,
            yearfrac_basis=int(yearfrac_basis),
            warnings=warnings,
            seen_warnings=seen_warnings,
        ))

    differences.sort(key=lambda x: (x.due_date, x.installment_number or 10**9))
    return MAJSResult(
        settings=settings, schedule=schedule, differences=differences, warnings=warnings,
        difference_correction_mode=difference_correction_mode,
        difference_interest_mode=difference_interest_mode,
        fixed_interest_start=fixed_interest_start,
        yearfrac_basis=int(yearfrac_basis),
        early_settlement=early_settlement,
    )


def majs_interest_mode_label(mode: str) -> str:
    return {
        'UM_PCT': '1% a.m. - pró-rata dias/30',
        'UM_PCT_YEARFRAC': '1% a.m. - fração de ano por convenção de dias',
        'UM_PCT_ATE_TL': '1% a.m. dias/30 até 29/08/2024 + Taxa Legal',
        'UM_PCT_YEARFRAC_ATE_TL': '1% a.m. por fração de ano até 29/08/2024 + Taxa Legal',
        'TAXA_LEGAL': 'Somente Taxa Legal a partir de 30/08/2024',
        'SEM_JUROS': 'Sem juros',
    }.get(mode, mode)


def majs_yearfrac_basis_label(basis: int) -> str:
    return {
        0: 'US/NASD 30/360',
        1: 'Real/Real',
        2: 'Real/360',
        3: 'Real/365',
        4: 'Europeu 30/360',
    }.get(int(basis), str(basis))


def majs_correction_mode_label(mode: str) -> str:
    if mode == 'INPC_IPCA_14905':
        return 'INPC até 08/2024 e IPCA a partir de 09/2024'
    if mode == SEM_CORRECAO:
        return 'Sem correção monetária'
    return f'{mode} em todo o período'


def build_majs_workbook(result: MAJSResult) -> bytes:
    """Exporta a memória MAJS para Excel com eventos, suspensões e cortes por falta de índice."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = 'RESUMO'
    teal = '005A70'; orange = 'F47C20'; white = 'FFFFFF'
    ws.append(['MOTOR DE CÁLCULOS - FUNCEF | RECÁLCULO MAJS'])
    ws['A1'].font = Font(bold=True, color=white, size=14)
    ws['A1'].fill = PatternFill('solid', fgColor=teal)
    ws.merge_cells('A1:D1')
    modality = 'Fixa — sem correção monetária' if result.settings.correction_index_code == SEM_CORRECAO else f'Variável — {result.settings.correction_index_code}'
    summary = [
        ('Valor original', result.settings.original_amount),
        ('Data do crédito', result.settings.credit_date),
        ('Taxa nominal anual', result.settings.monthly_interest_rate * Decimal('12')),
        ('Prazo em prestações', result.settings.term_months),
        ('Modalidade', modality),
        ('Data-base solicitada', result.settings.base_date),
        ('Última evolução MAJS', result.last_evolution_date),
        ('Quitação antecipada', 'Sim' if result.was_settled_early else 'Não'),
        ('Data da quitação', result.early_settlement.settlement_date if result.was_settled_early else '-'),
        ('Forma da quitação', ('Novação' if result.early_settlement and result.early_settlement.novation else 'Pagamento pelo participante') if result.was_settled_early else '-'),
        ('Saldo MAJS na quitação', result.early_settlement_balance if result.was_settled_early else Decimal('0')),
        ('Valor pago na quitação', result.early_settlement.amount_paid if result.was_settled_early and result.early_settlement and not result.early_settlement.novation else Decimal('0')),
        ('Correção das diferenças', majs_correction_mode_label(result.difference_correction_mode)),
        ('Juros sobre as diferenças', majs_interest_mode_label(result.difference_interest_mode)),
        ('Início dos juros de mora', result.fixed_interest_start or 'Data de origem de cada diferença'),
        ('Convenção de contagem de dias', majs_yearfrac_basis_label(result.yearfrac_basis) if 'YEARFRAC' in result.difference_interest_mode else 'Não se aplica'),
        ('Total devido até a data-base', result.due_total),
        ('Total pago / crédito reconciliado', result.paid_total),
        ('Diferença pago - devido', result.difference_total),
        ('Diferença atualizada', result.updated_difference_total),
        ('Saldo teórico a vencer', result.balance_to_mature),
    ]
    for label, value in summary:
        ws.append([label, value])
    for cell in ws['A']:
        cell.font = Font(bold=True)
    ws.column_dimensions['A'].width = 38; ws.column_dimensions['B'].width = 28
    for row in range(2, ws.max_row + 1):
        value = ws.cell(row, 2).value
        if isinstance(value, Decimal):
            ws.cell(row, 2).value = float(value)
            ws.cell(row, 2).number_format = '#,##0.00;[Red](#,##0.00);-'
        elif isinstance(value, date):
            ws.cell(row, 2).number_format = 'dd/mm/yyyy'

    evo = wb.create_sheet('EVOLUCAO_MAJS')
    headers = ['Evento','Prestação','Data','Saldo inicial','Posição fator','Fator financeiro','Soma fatores','Prestação devida / saldo quitação','Juros','Amortização','Ref. índice','Correção aplicada %','Amortização extra aplicada','Valor pago na quitação','Excesso amortização','Saldo final','Observação']
    evo.append(headers)
    for c in evo[1]:
        c.fill = PatternFill('solid', fgColor=teal); c.font=Font(color=white,bold=True); c.alignment=Alignment(horizontal='center')
    for r in result.schedule:
        evo.append([
            {'SUSPENSAO':'Suspensão','AMORTIZACAO_EXTRA':'Amortização extraordinária','PRESTACAO':'Prestação','QUITACAO_ANTECIPADA':'Quitação antecipada'}.get(r.event_type, r.event_type),
            r.installment_number, r.due_date, float(r.opening_balance), r.factor_position,
            float(r.financial_factor), float(r.remaining_factor_sum), float(r.installment_due),
            float(r.interest_component), float(r.amortization_component), r.correction_reference,
            float(r.correction_rate), float(r.extra_amortization), float(r.settlement_payment), float(r.excess_extra_amortization),
            float(r.closing_balance), r.event_note,
        ])
    for col in ['D','H','I','J','M','N','O','P']:
        for cell in evo[col][1:]: cell.number_format='#,##0.00;[Red](#,##0.00);-'
    for cell in evo['L'][1:]: cell.number_format='0.0000%'
    for cell in evo['C'][1:]: cell.number_format='dd/mm/yyyy'
    evo.freeze_panes='A2'
    for idx, width in enumerate([18,11,13,16,13,16,14,22,14,14,14,12,21,20,18,16,42], start=1):
        evo.column_dimensions[get_column_letter(idx)].width = width

    dif = wb.create_sheet('DIFERENCAS')
    headers2=['Origem','Prestação','Data de origem','Devido','Pago/crédito','Diferença (pago - devido)','Fator correção','Diferença corrigida','Início dos juros de mora','Juros acumulados %','Juros R$','Diferença atualizada','Atualizado até','Observação']
    dif.append(headers2)
    for c in dif[1]:
        c.fill=PatternFill('solid',fgColor=teal); c.font=Font(color=white,bold=True); c.alignment=Alignment(horizontal='center')
    for r in result.differences:
        dif.append([
            r.origin, r.installment_number, r.due_date, float(r.amount_due), float(r.amount_paid),
            float(r.difference_paid_minus_due), float(r.correction_factor), float(r.corrected_difference),
            r.interest_start_date, float(r.accumulated_interest_rate), float(r.interest_amount), float(r.updated_difference),
            r.updated_through_date, r.note,
        ])
    for col in ['D','E','F','H','K','L']:
        for cell in dif[col][1:]: cell.number_format='#,##0.00;[Red](#,##0.00);-'
    for cell in dif['J'][1:]: cell.number_format='0.0000%'
    for col in ['C','I','M']:
        for cell in dif[col][1:]: cell.number_format='dd/mm/yyyy'
    dif.freeze_panes='A2'
    for idx, width in enumerate([34,11,15,14,14,20,15,18,19,19,14,18,15,38], start=1):
        dif.column_dimensions[get_column_letter(idx)].width = width

    avisos = wb.create_sheet('AVISOS')
    avisos.append(['Avisos / premissas'])
    avisos['A1'].fill=PatternFill('solid',fgColor=orange); avisos['A1'].font=Font(color=white,bold=True)
    avisos.append(['Pagamentos ordinários são comparativos e não alteram o saldo teórico MAJS.'])
    avisos.append(['Se a amortização extraordinária ocorrer entre vencimentos, aplica-se primeiro somente a correção proporcional aos dias transcorridos, por equivalência diária; no vencimento aplica-se a parcela remanescente do índice. Depois da correção proporcional, a amortização reduz o saldo e reinicia a sequência do fator financeiro.'])
    avisos.append(['Se a amortização extraordinária superar o saldo, o excedente é lançado como diferença credora e não é descartado.'])
    avisos.append(['Na suspensão, o saldo é corrigido, mas não há prestação, juros ou amortização contratual; a quantidade de prestações do prazo é preservada.'])
    avisos.append(['Na modalidade fixa, a correção monetária do saldo MAJS é zero.'])
    avisos.append(['Na falta de índice, o cálculo é parcial e não projeta competências inexistentes.'])
    avisos.append(['Na quitação antecipada, o fluxo contratual é encerrado na data informada. Em novação, o saldo de quitação não é tratado como pagamento do participante nem compõe diferenças; sem novação, a diferença entre saldo teórico e valor efetivamente pago compõe a tabela de diferenças.'])
    if result.difference_interest_mode == 'UM_PCT_YEARFRAC':
        avisos.append([f'Juros de mora: 12% a.a. em regime simples, proporcionalmente à fração de ano apurada pela convenção {majs_yearfrac_basis_label(result.yearfrac_basis)}.'])
    elif result.difference_interest_mode == 'UM_PCT_YEARFRAC_ATE_TL':
        avisos.append([f'Até 29/08/2024: 12% a.a. em regime simples, proporcionalmente à fração de ano pela convenção {majs_yearfrac_basis_label(result.yearfrac_basis)}; a partir de 30/08/2024: Taxa Legal.'])
    for warning in result.warnings: avisos.append([warning])
    avisos.column_dimensions['A'].width=140

    bio=BytesIO(); wb.save(bio); return bio.getvalue()
