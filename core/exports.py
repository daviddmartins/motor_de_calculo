from __future__ import annotations

from dataclasses import asdict
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .models import ContractSettings, EvolutionResult, PRICE_MODALITIES


PETROLEUM = "005A70"
ORANGE = "F28C28"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D5E1E5")
WEEKDAYS_PT = [
    "Segunda-feira", "Terça-feira", "Quarta-feira", "Quinta-feira",
    "Sexta-feira", "Sábado", "Domingo",
]


def _style_sheet(ws, header_row: int = 1) -> None:
    for cell in ws[header_row]:
        cell.fill = PatternFill("solid", fgColor=PETROLEUM)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top")
    for col in range(1, ws.max_column + 1):
        width = 12
        for cell in ws[get_column_letter(col)]:
            width = min(max(width, len(str(cell.value or "")) + 2), 32)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = ws.dimensions


def _money_format(places: int) -> str:
    decimals = "." + ("0" * places) if places else ""
    return f'R$ #,##0{decimals};[Red]-R$ #,##0{decimals}'


def _percent_format(places: int) -> str:
    decimals = "." + ("0" * places) if places else ""
    return f"0{decimals}%"


def build_evolution_workbook(
    settings: ContractSettings,
    modality_name: str,
    result: EvolutionResult,
) -> bytes:
    money_format = _money_format(settings.rounding.money_places)
    percent_format = _percent_format(settings.rounding.percentage_display_places)

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "MEMÓRIA DE CÁLCULO — VERSÃO DE VALIDAÇÃO"
    ws["A1"].fill = PatternFill("solid", fgColor=PETROLEUM)
    ws["A1"].font = Font(color=WHITE, bold=True, size=14)
    ws.merge_cells("A1:D1")
    rows = [
        ("Modalidade", modality_name),
        ("Código interno", settings.modality_code),
        ("Saldo inicial", float(settings.initial_balance)),
        ("Data do crédito", settings.credit_date),
        ("Taxa anual", float(settings.annual_interest_rate)),
        ("Taxa mensal equivalente", float(result.monthly_interest_rate)),
        ("Prazo", settings.term),
        ("Início da competência", settings.competence_start_day),
        ("Defasagem do índice", settings.index_lag_months),
        ("Tipo da fórmula Price", settings.payment_timing if settings.modality_code in PRICE_MODALITIES else "Não se aplica"),
        ("Prestação Price calculada", float(result.fixed_payment) if result.fixed_payment else None),
        ("Casas exibidas nos percentuais", settings.rounding.percentage_display_places),
        ("Casas exibidas nos valores", settings.rounding.money_places),
        ("Status do cálculo", "Completo" if result.is_complete else "Parcial"),
        ("Prestações solicitadas", result.requested_installments or settings.term),
        ("Prestações concluídas", result.processed_installments),
        ("Última data calculada", result.last_calculated_date),
        ("Índice ausente", result.missing_index_reference or ""),
        ("Nota técnica", result.status_note or "Cálculo concluído para todo o período solicitado."),
    ]
    for i, (label, value) in enumerate(rows, start=3):
        ws.cell(i, 1, label).font = Font(bold=True, color=PETROLEUM)
        ws.cell(i, 2, value)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80
    for row_idx in range(3, 3 + len(rows)):
        label = ws.cell(row_idx, 1).value
        if label in {"Saldo inicial", "Prestação Price calculada"}:
            ws.cell(row_idx, 2).number_format = money_format
        elif label in {"Taxa anual", "Taxa mensal equivalente"}:
            ws.cell(row_idx, 2).number_format = percent_format
        elif label in {"Data do crédito", "Última data calculada"}:
            ws.cell(row_idx, 2).number_format = "dd/mm/yyyy"
        ws.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")

    inst = wb.create_sheet("Prestacoes")
    headers = [
        "Prestação", "Data de referência", "Vencimento operacional", "Competência",
        "Índice de referência", "Saldo inicial do período", "Saldo antes da prestação",
        "Base da amortização", "Correção monetária", "Juros do período",
        "Amortização regular", "Amortização extraordinária", "Valor da prestação",
        "Saldo final", "Parcelas remanescentes",
    ]
    inst.append(headers)
    for row in result.installments:
        inst.append([
            row.installment_number, row.reference_due_date, row.operational_due_date,
            row.competence, row.index_reference, float(row.opening_balance),
            float(row.balance_before_installment), float(row.amortization_base),
            float(row.correction_amount), float(row.interest_amount),
            float(row.regular_amortization), float(row.extraordinary_amortization),
            float(row.installment_amount), float(row.closing_balance), row.remaining_installments,
        ])
    _style_sheet(inst)
    for col in (2, 3):
        for cell in inst[get_column_letter(col)][1:]:
            cell.number_format = "dd/mm/yyyy"
    for col in range(6, 15):
        for cell in inst[get_column_letter(col)][1:]:
            cell.number_format = money_format

    daily = wb.create_sheet("Memoria_Diaria")
    daily_headers = [
        "Prestação", "Dia da semana", "Feriado", "Data", "Competência",
        "Índice referência", "Dia útil", "Saldo devedor inicial", "Juros (%)",
        "Juros (R$)", "S.D. com juros", "C.M. (%)", "C.M. (R$)",
        "Juros + C.M. (R$)", "Saldo devedor (J+CM)", "Amortização extraordinária",
        "Juros acumulados", "Amortização regular", "Prestação",
        "Saldo devedor finalizado do dia",
    ]
    daily.append(daily_headers)
    for row in result.daily_rows:
        finalized = (
            row.closing_balance_after_installment
            if row.closing_balance_after_installment is not None
            else row.balance_before_installment
        )
        daily.append([
            row.installment_number,
            WEEKDAYS_PT[row.day.weekday()],
            "Sim" if row.is_holiday else "",
            row.day,
            row.competence,
            row.index_reference,
            "Sim" if row.is_business_day else "Não",
            float(row.opening_balance),
            float(row.daily_interest_rate),
            float(row.interest_amount),
            float(row.balance_after_interest),
            float(row.daily_correction_rate),
            float(row.correction_amount),
            float(row.interest_amount + row.correction_amount),
            float(row.balance_after_correction),
            float(row.extraordinary_amortization),
            float(row.accumulated_interest),
            float(row.regular_amortization),
            float(row.installment_amount),
            float(finalized),
        ])
    _style_sheet(daily)
    for cell in daily["D"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for col in (9, 12):
        for cell in daily[get_column_letter(col)][1:]:
            cell.number_format = percent_format
    for col in (8, 10, 11, 13, 14, 15, 16, 17, 18, 19, 20):
        for cell in daily[get_column_letter(col)][1:]:
            cell.number_format = money_format

    params = wb.create_sheet("Parametros")
    params.append(["Parâmetro", "Valor"])
    for key, value in asdict(settings).items():
        params.append([key, str(value)])
    _style_sheet(params)

    warnings = wb.create_sheet("Avisos")
    warnings.append(["Aviso"])
    for warning in result.warnings or ["Nenhum aviso gerado."]:
        warnings.append([warning])
    _style_sheet(warnings)
    warnings.column_dimensions["A"].width = 100
    warnings["A1"].fill = PatternFill("solid", fgColor=ORANGE)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
