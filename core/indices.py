from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import BinaryIO
import csv

from openpyxl import load_workbook


@dataclass(frozen=True)
class IndexSeries:
    code: str
    values: dict[str, Decimal]
    source: str = ""

    def get(self, competence: str) -> Decimal:
        try:
            return self.values[competence]
        except KeyError as exc:
            raise KeyError(f"Índice {self.code} ausente para a competência {competence}.") from exc


@dataclass
class AppConfiguration:
    modality_names: dict[str, str]
    indices: dict[str, IndexSeries]
    holidays: set[date]
    parameters: dict[str, str]


def _excel_date(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return date(1899, 12, 30) + timedelta(days=int(value))
    return None


def _competence_key(value) -> str | None:
    d = _excel_date(value)
    if d:
        return f"{d.year:04d}-{d.month:02d}"
    if value is None:
        return None
    text = str(value).strip()
    for fmt in ("%Y-%m", "%m/%Y", "%Y/%m", "%d/%m/%Y"):
        try:
            parsed = datetime.strptime(text, fmt)
            return f"{parsed.year:04d}-{parsed.month:02d}"
        except ValueError:
            pass
    return None


def _decimal_percent(value) -> Decimal | None:
    if value is None or value == "":
        return None
    return Decimal(str(value).replace("%", "").replace(",", ".")) / Decimal("100")


def _load_series_from_sheet(ws, code: str) -> IndexSeries:
    header_row = None
    competence_col = None
    percent_col = None
    max_scan = min(ws.max_row, 20)
    for row in range(1, max_scan + 1):
        values = [str(ws.cell(row, col).value or "").strip().lower() for col in range(1, min(ws.max_column, 12) + 1)]
        for idx, text in enumerate(values, start=1):
            if "compet" in text:
                competence_col = idx
            if "percentual" in text or "% no mês" in text or "indice" == text or "índice" == text:
                percent_col = idx
        if competence_col and percent_col:
            header_row = row
            break
    if not header_row:
        raise ValueError(f"Não foi possível identificar as colunas da aba {ws.title}.")

    values: dict[str, Decimal] = {}
    for row in range(header_row + 1, ws.max_row + 1):
        key = _competence_key(ws.cell(row, competence_col).value)
        rate = _decimal_percent(ws.cell(row, percent_col).value)
        if key and rate is not None:
            values[key] = rate
    if not values:
        raise ValueError(f"A aba {ws.title} não contém índices válidos.")
    return IndexSeries(code=code, values=values, source=ws.title)



def _load_series_from_csv(path: Path, code: str) -> IndexSeries:
    values: dict[str, Decimal] = {}
    source_name = path.name
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            key = str(row.get("competencia") or row.get("Competência") or "").strip()
            raw = row.get("percentual_mensal_pct") or row.get("Percentual mensal (%)")
            if not key or raw in (None, ""):
                continue
            try:
                values[key] = Decimal(str(raw).replace("%", "").replace(",", ".")) / Decimal("100")
            except Exception:
                continue
            src = str(row.get("fonte") or row.get("Fonte") or "").strip()
            if src:
                source_name = src
    if not values:
        raise ValueError(f"O arquivo {path.name} não contém índices válidos.")
    return IndexSeries(code=code, values=values, source=source_name)


def _load_external_index_series(indices: dict[str, IndexSeries]) -> None:
    data_dir = Path(__file__).resolve().parent.parent / "data" / "indices"
    files = {
        "IGPM": "igpm_sgs_28655.csv",
        "TR": "tr_sgs_7811.csv",
        "TAXA_LEGAL": "taxa_legal_sgs_29543.csv",
    }
    for code, filename in files.items():
        path = data_dir / filename
        if not path.exists():
            continue
        try:
            indices[code] = _load_series_from_csv(path, code)
        except ValueError:
            continue

def load_configuration(source: str | Path | bytes | BinaryIO) -> AppConfiguration:
    if isinstance(source, (str, Path)):
        wb = load_workbook(source, data_only=True, read_only=False)
    elif isinstance(source, bytes):
        wb = load_workbook(BytesIO(source), data_only=True, read_only=False)
    else:
        wb = load_workbook(source, data_only=True, read_only=False)

    modality_names: dict[str, str] = {
        "MOD_001": "C Variável",
        "MOD_002": "C Fixo",
        "MOD_004": "Novo Credinâmico Fixo",
        "MOD_005": "Credinâmico Fixo",
        "MOD_006": "Novo Credinâmico Variável",
    }
    parameters: dict[str, str] = {}
    holidays: set[date] = set()
    indices: dict[str, IndexSeries] = {}

    if "MODALIDADES" in wb.sheetnames:
        ws = wb["MODALIDADES"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] and row[1]:
                modality_names[str(row[0]).strip()] = str(row[1]).strip()
    elif "02_PARAM_MODALIDADES" in wb.sheetnames:
        ws = wb["02_PARAM_MODALIDADES"]
        for row in ws.iter_rows(min_row=5, values_only=True):
            if row[0] and row[1]:
                modality_names[str(row[0]).strip()] = str(row[1]).strip()

    if "PARAMETROS" in wb.sheetnames:
        ws = wb["PARAMETROS"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is not None:
                parameters[str(row[0]).strip()] = str(row[1]).strip() if row[1] is not None else ""

    if "FERIADOS" in wb.sheetnames:
        ws = wb["FERIADOS"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            d = _excel_date(row[0])
            if d:
                holidays.add(d)

    candidate_sheets = {
        "INPC": ["INPC", "04_INPC"],
        "IPCA": ["IPCA", "05_IPCA"],
        "SELIC_M": ["SELIC_M", "06_SELIC_MENSAL"],
    }
    for code, names in candidate_sheets.items():
        for name in names:
            if name in wb.sheetnames:
                try:
                    indices[code] = _load_series_from_sheet(wb[name], code)
                except ValueError:
                    pass
                break

    wb.close()
    _load_external_index_series(indices)
    return AppConfiguration(modality_names=modality_names, indices=indices, holidays=holidays, parameters=parameters)
