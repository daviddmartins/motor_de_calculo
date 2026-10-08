from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook


@dataclass(frozen=True)
class RequestIdentity:
    email: str = ""
    preferred_username: str = ""
    user_id: str = ""

    @property
    def fallback_name(self) -> str:
        return self.preferred_username or self.email or self.user_id or "Usuário logado"


@dataclass(frozen=True)
class AccessEntry:
    email: str
    name: str
    profile: str
    active: bool

    @property
    def is_admin(self) -> bool:
        return _norm(self.profile) in {"administrador", "admin"}


@dataclass(frozen=True)
class AccessRegistry:
    entries: tuple[AccessEntry, ...]
    source_path: str = ""

    @property
    def enabled(self) -> bool:
        # O controle só é ativado quando existe ao menos um Login_Email preenchido.
        return any(x.email.strip() for x in self.entries)

    @property
    def active_entries(self) -> tuple[AccessEntry, ...]:
        return tuple(x for x in self.entries if x.active and x.email.strip())

    @property
    def active_names(self) -> list[str]:
        # Lista usada nos combos de validação. Evita nomes duplicados.
        names = {x.name.strip() for x in self.active_entries if x.name.strip()}
        return sorted(names, key=lambda x: x.casefold())

    @property
    def email_count(self) -> int:
        return sum(1 for x in self.entries if x.email.strip())

    def find(self, email: str) -> AccessEntry | None:
        key = _email_key(email)
        if not key:
            return None
        for entry in self.entries:
            if _email_key(entry.email) == key:
                return entry
        return None


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def _email_key(value: object) -> str:
    text = str(value or "").strip().casefold()
    # Remove prefixos comuns quando a célula foi colada como mailto:usuario@dominio.
    if text.startswith("mailto:"):
        text = text[7:]
    return text


def _active(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return _norm(value) in {"true", "verdadeiro", "sim", "yes", "1", "ativo", "ativa", "x"}


def detect_request_identity(headers: dict | None) -> RequestIdentity:
    normalized = {
        str(k).strip().casefold(): str(v).strip()
        for k, v in (headers or {}).items()
        if v is not None
    }
    return RequestIdentity(
        email=normalized.get("x-forwarded-email", ""),
        preferred_username=normalized.get("x-forwarded-preferred-username", ""),
        user_id=normalized.get("x-forwarded-user", ""),
    )


def load_access_registry(path: str | Path) -> AccessRegistry:
    path = Path(path)
    if not path.exists() or not path.is_file():
        return AccessRegistry(entries=tuple(), source_path=str(path))

    wb = load_workbook(path, data_only=True, read_only=True)
    sheet = wb["Controle_De_Acesso"] if "Controle_De_Acesso" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return AccessRegistry(entries=tuple(), source_path=str(path))

    header = {_norm(value): idx for idx, value in enumerate(rows[0]) if value is not None}

    def col(*aliases: str) -> int | None:
        for alias in aliases:
            key = _norm(alias)
            if key in header:
                return header[key]
        return None

    email_col = col("Login_Email", "Email", "E-mail", "Login", "Login Email")
    name_col = col("Nome", "Nome completo", "Nome_Completo")
    # O modelo recebido contém "Pefil"; a versão oficial pode usar "Perfil".
    profile_col = col("Perfil", "Pefil", "Profile")
    active_col = col("Ativo", "Ativa", "Status")

    if email_col is None or name_col is None or profile_col is None or active_col is None:
        raise ValueError(
            f"A planilha {path.name} deve conter as colunas Login_Email, Nome, Perfil/Pefil e Ativo."
        )

    entries: list[AccessEntry] = []
    for raw in rows[1:]:
        email = str(raw[email_col] or "").strip() if email_col < len(raw) else ""
        name = str(raw[name_col] or "").strip() if name_col < len(raw) else ""
        profile = str(raw[profile_col] or "Usuário").strip() if profile_col < len(raw) else "Usuário"
        active = _active(raw[active_col] if active_col < len(raw) else False)
        if email or name:
            entries.append(
                AccessEntry(
                    email=email,
                    name=name,
                    profile=profile or "Usuário",
                    active=active,
                )
            )
    return AccessRegistry(entries=tuple(entries), source_path=str(path))


def _candidate_access_files(app_dir: Path, preferred_path: Path | None = None) -> list[Path]:
    """Localiza a planilha mesmo quando o arquivo oficial foi importado com outro nome.

    Aceita o arquivo na raiz do app ou em ``config`` e prioriza nomes conhecidos. Também
    considera qualquer .xlsx cujo nome contenha "acess". Arquivos temporários do Excel
    (``~$``) são ignorados.
    """
    app_dir = Path(app_dir)
    config_dir = app_dir / "config"
    names = [
        "parametros_acesso_motor_calculos.xlsx",
        "Parametros_Acesso_Motor_Calculos.xlsx",
        "Parametros_Acessop_Motor_Calculos.xlsx",  # nome do modelo recebido
        "parametros_acessop_motor_calculos.xlsx",
    ]

    ordered: list[Path] = []
    if preferred_path:
        ordered.append(Path(preferred_path))
    for folder in (config_dir, app_dir):
        ordered.extend(folder / name for name in names)
        if folder.exists():
            ordered.extend(
                p for p in sorted(folder.glob("*.xlsx"))
                if not p.name.startswith("~$") and "acess" in _norm(p.stem)
            )

    seen: set[str] = set()
    result: list[Path] = []
    for path in ordered:
        key = str(path.resolve()) if path.exists() else str(path.absolute())
        if key not in seen:
            seen.add(key)
            result.append(path)
    return result


def discover_access_registry(app_dir: str | Path, preferred_path: str | Path | None = None) -> AccessRegistry:
    """Escolhe automaticamente a planilha de acesso efetivamente preenchida.

    Esse comportamento evita que o arquivo-modelo vazio incluído no deploy tenha
    prioridade sobre uma planilha oficial que tenha sido importada na raiz ou em config.
    Se houver mais de uma planilha válida, vence a que possuir mais e-mails ativos e,
    depois, mais e-mails preenchidos.
    """
    candidates = _candidate_access_files(Path(app_dir), Path(preferred_path) if preferred_path else None)
    parsed: list[tuple[tuple[int, int, int], int, AccessRegistry]] = []

    for order, path in enumerate(candidates):
        if not path.exists() or not path.is_file():
            continue
        try:
            registry = load_access_registry(path)
        except Exception:
            # Um .xlsx de outro propósito não deve derrubar o app durante a descoberta.
            continue
        score = (
            len(registry.active_entries),
            registry.email_count,
            1 if registry.enabled else 0,
        )
        parsed.append((score, -order, registry))

    if not parsed:
        fallback = Path(preferred_path) if preferred_path else Path(app_dir) / "config" / "parametros_acesso_motor_calculos.xlsx"
        return AccessRegistry(entries=tuple(), source_path=str(fallback))

    parsed.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return parsed[0][2]
