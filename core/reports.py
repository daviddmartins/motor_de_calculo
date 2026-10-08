from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from datetime import date, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import re
import unicodedata

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph

from .calendar_rules import calendar_days_in_day_competence
from .models import ContractSettings, EvolutionResult


MONTHS_PT = [
    "Jan", "Fev", "Mar", "Abr", "Mai", "Jun",
    "Jul", "Ago", "Set", "Out", "Nov", "Dez",
]

# Identidade visual: azul-petróleo como cor principal, azul FUNCEF e laranja
# apenas como acento. Sem gradientes.
PETROLEUM = "0B4F6C"
FUNCEF_BLUE = "155AA8"
FUNCEF_ORANGE = "F58220"
LIGHT_BLUE = "EAF3F7"
LIGHTER_BLUE = "F5F9FB"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "CBD5E1"
DARK_TEXT = "26333D"
WHITE = "FFFFFF"


@dataclass(frozen=True)
class OpinionData:
    contract_number: str
    participant_name: str
    registration: str
    cpf: str
    request_date: date
    margin_amount: Decimal
    fgqc_concession: Decimal
    iof_amount: Decimal
    administrative_fee: Decimal
    settled_loan_balance: Decimal
    net_amount: Decimal
    first_fgqc: Decimal = Decimal("0")
    second_fgqc: Decimal = Decimal("0")


def _number_br(value: Decimal | int | float | str, places: int = 2) -> str:
    number = Decimal(str(value))
    text = f"{number:,.{places}f}"
    return text.replace(",", "X").replace(".", ",").replace("X", ".")


def _money_br(value: Decimal | int | float | str, places: int = 2) -> str:
    return f"R$ {_number_br(value, places)}"


def _percent_br(rate: Decimal, places: int = 4) -> str:
    return f"{_number_br(rate * Decimal('100'), places)}%"


def _date_br(value: date) -> str:
    return value.strftime("%d/%m/%Y")


def _month_year(value: date) -> str:
    return f"{MONTHS_PT[value.month - 1]}/{value.year}"


def _copy_first_run_format(paragraph: Paragraph):
    if not paragraph.runs:
        return None
    rpr = paragraph.runs[0]._r.rPr
    return deepcopy(rpr) if rpr is not None else None


def _set_paragraph(paragraph: Paragraph, text: str) -> None:
    rpr = _copy_first_run_format(paragraph)
    paragraph.clear()
    run = paragraph.add_run(text)
    if rpr is not None:
        run._r.insert(0, rpr)


def _set_cell(cell: _Cell, text: str) -> None:
    paragraph = cell.paragraphs[0]
    _set_paragraph(paragraph, text)
    for extra in cell.paragraphs[1:]:
        _set_paragraph(extra, "")


def _set_table_widths(table, widths: list[float]) -> None:
    table.autofit = False
    twips = [Inches(width).twips for width in widths]
    tblw = table._tbl.tblPr.first_child_found_in("w:tblW")
    if tblw is not None:
        tblw.set(qn("w:type"), "dxa")
        tblw.set(qn("w:w"), str(sum(twips)))
    grid_cols = table._tbl.tblGrid.gridCol_lst
    for idx, width_twips in enumerate(twips):
        if idx < len(grid_cols):
            grid_cols[idx].set(qn("w:w"), str(width_twips))
        table.columns[idx].width = Inches(widths[idx])
    for row in table.rows:
        for idx, width_twips in enumerate(twips):
            cell = row.cells[idx]
            cell.width = Inches(widths[idx])
            tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tcw.set(qn("w:type"), "dxa")
            tcw.set(qn("w:w"), str(width_twips))


def _set_cell_font_size(cell: _Cell, points: float) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(points)


def _safe_filename(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    normalized = re.sub(r"[^A-Za-z0-9._-]+", "_", normalized).strip("_")
    return normalized or "parecer"


def opinion_filename(modality_name: str, contract_number: str) -> str:
    return _safe_filename(f"Parecer_{modality_name}_{contract_number}") + ".docx"


def _math_nodes(paragraph: Paragraph):
    return paragraph._p.xpath(".//m:t")


def _set_math_texts(paragraph: Paragraph, values: list[str]) -> None:
    nodes = _math_nodes(paragraph)
    if len(nodes) != len(values):
        raise ValueError(
            f"A estrutura matemática do modelo foi alterada: esperados {len(values)} campos, encontrados {len(nodes)}."
        )
    for node, value in zip(nodes, values):
        node.text = value


def _period_rows(result: EvolutionResult, installment_number: int):
    return [row for row in result.daily_rows if row.installment_number == installment_number]


def _period_day_count(result: EvolutionResult, installment_number: int) -> int:
    return len(_period_rows(result, installment_number))


def _single_competence_denominator(
    result: EvolutionResult,
    installment_number: int,
    settings: ContractSettings,
) -> int | None:
    rows = _period_rows(result, installment_number)
    if not rows or len({row.competence for row in rows}) != 1:
        return None
    return calendar_days_in_day_competence(
        rows[0].day,
        settings.competence_start_day,
        settings.due_day,
    )


def _format_contract_table(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
    modality_name: str,
    data: OpinionData,
) -> None:
    table = doc.tables[0]
    _set_table_widths(table, [3.75, 2.55])
    first_due = result.installments[0].reference_due_date
    system_name = (
        "Sistema de Amortização Constante – SAC"
        if settings.modality_code == "MOD_001"
        else "Sistema Francês de Amortização – Tabela Price"
    )
    values: list[str] = [
        _date_br(data.request_date),
        _date_br(settings.credit_date),
        _money_br(data.margin_amount, settings.rounding.money_places),
        _money_br(settings.initial_balance, settings.rounding.money_places),
        _money_br(data.fgqc_concession, settings.rounding.money_places),
        _money_br(data.iof_amount, settings.rounding.money_places),
        _money_br(data.administrative_fee, settings.rounding.money_places),
        _money_br(data.settled_loan_balance, settings.rounding.money_places),
        _money_br(data.net_amount, settings.rounding.money_places),
        _date_br(first_due),
        system_name,
    ]
    if settings.modality_code == "MOD_002":
        values.append(_money_br(result.fixed_payment or Decimal("0"), settings.rounding.money_places))
    values.extend([
        _percent_br(settings.annual_interest_rate, settings.rounding.percentage_display_places),
        str(settings.term),
    ])
    if len(table.rows) != len(values):
        raise ValueError("O modelo Word não possui a tabela de parâmetros esperada para esta modalidade.")
    for row, value in zip(table.rows, values):
        _set_cell(row.cells[1], value)
        _set_cell_font_size(row.cells[1], 9)


def _fill_variable_summary_table(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
    data: OpinionData,
) -> None:
    table = doc.tables[1]
    first = result.installments[0]
    second = result.installments[1] if len(result.installments) > 1 else None
    places = settings.rounding.money_places
    initial = [
        _month_year(settings.credit_date), "Valor Contratado", "", "-", "-", "-", "-",
        _number_br(settings.initial_balance, places), "-",
    ]
    rows = [initial]
    for item, fgqc in ((first, data.first_fgqc), (second, data.second_fgqc)):
        if item is None:
            rows.append(["", "", "", "", "", "", "", "", ""])
            continue
        rows.append([
            _month_year(item.reference_due_date),
            f"Prestação - {item.installment_number:03d}/{item.remaining_installments:03d}",
            _date_br(item.reference_due_date),
            _number_br(item.installment_amount, places),
            _number_br(item.regular_amortization, places),
            _number_br(item.interest_amount, places),
            _number_br(fgqc, places) if fgqc else "-",
            _number_br(item.closing_balance, places),
            _number_br(item.correction_amount, places),
        ])
    for row_idx, values in enumerate(rows, start=1):
        for col_idx, value in enumerate(values):
            _set_cell(table.rows[row_idx].cells[col_idx], value)


def _fill_fixed_summary_tables(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
    data: OpinionData,
) -> None:
    places = settings.rounding.money_places
    pct_places = settings.rounding.percentage_display_places
    rates_table = doc.tables[1]
    _set_cell(rates_table.rows[0].cells[1], _percent_br(settings.annual_interest_rate, pct_places))
    _set_cell(rates_table.rows[1].cells[1], _percent_br(result.monthly_interest_rate, pct_places))

    table = doc.tables[2]
    first = result.installments[0]
    second = result.installments[1] if len(result.installments) > 1 else None
    rows = [[
        _month_year(settings.credit_date), "Valor Contratado", "", "-", "-", "-", "-",
        _number_br(settings.initial_balance, places),
    ]]
    for item, fgqc in ((first, data.first_fgqc), (second, data.second_fgqc)):
        if item is None:
            rows.append(["", "", "", "", "", "", "", ""])
            continue
        rows.append([
            _month_year(item.reference_due_date),
            f"Prestação - {item.installment_number:03d}/{item.remaining_installments:03d}",
            _date_br(item.reference_due_date),
            _number_br(item.installment_amount, places),
            _number_br(item.regular_amortization, places),
            _number_br(item.interest_amount, places),
            _number_br(fgqc, places) if fgqc else "-",
            _number_br(item.closing_balance, places),
        ])
    for row_idx, values in enumerate(rows, start=1):
        for col_idx, value in enumerate(values):
            _set_cell(table.rows[row_idx].cells[col_idx], value)


def _fill_variable_narrative(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
) -> None:
    first = result.installments[0]
    second = result.installments[1] if len(result.installments) > 1 else None
    places = settings.rounding.money_places
    pct_places = settings.rounding.percentage_display_places
    paragraphs = doc.paragraphs

    _set_paragraph(
        paragraphs[48],
        "Na data de referência da 1ª prestação, o saldo acumulado com juros e correção monetária "
        f"totalizou {_money_br(first.balance_before_installment, places)}. Deduzidos os juros gerados "
        f"no período, a base de amortização foi de {_money_br(first.amortization_base, places)}. "
        f"Dividindo-se essa base pelas {settings.term} prestações, a amortização apurada foi de "
        f"{_money_br(first.regular_amortization, places)}.",
    )
    _set_paragraph(
        paragraphs[49],
        f"A = {_number_br(first.amortization_base, places)}/{settings.term} = "
        f"{_money_br(first.regular_amortization, places)}",
    )
    first_days = _period_day_count(result, 1)
    _set_paragraph(
        paragraphs[53],
        f"A taxa de juros do contrato é de {_percent_br(settings.annual_interest_rate, pct_places)} ao ano, "
        f"equivalente a {_percent_br(result.monthly_interest_rate, pct_places)} ao mês. A taxa mensal é "
        "transformada em taxas diárias equivalentes, de acordo com a quantidade de dias corridos de cada "
        f"competência, e aplicada diariamente sobre o saldo devedor. No período de carência de {first_days} "
        f"dias, o somatório dos juros diários totalizou {_money_br(first.interest_amount, places)}.",
    )
    _set_paragraph(
        paragraphs[58],
        f"P = {_money_br(first.regular_amortization, places)} + {_money_br(first.interest_amount, places)} "
        f"= {_money_br(first.installment_amount, places)}",
    )
    _set_paragraph(
        paragraphs[59],
        f"Dessa forma, a primeira prestação do contrato foi de {_money_br(first.installment_amount, places)}, "
        f"composta pela amortização de {_money_br(first.regular_amortization, places)} e pelos juros de "
        f"{_money_br(first.interest_amount, places)}.",
    )
    if second is not None:
        remaining = settings.term - 1
        _set_paragraph(
            paragraphs[63],
            "Na data de referência da 2ª prestação, o saldo acumulado com juros e correção monetária "
            f"totalizou {_money_br(second.balance_before_installment, places)}. Deduzidos os juros do período, "
            f"a base de amortização foi de {_money_br(second.amortization_base, places)}. Dividindo-se essa "
            f"base pelas {remaining} prestações remanescentes, a amortização apurada foi de "
            f"{_money_br(second.regular_amortization, places)}.",
        )
        _set_paragraph(
            paragraphs[64],
            f"A = {_number_br(second.amortization_base, places)}/{remaining} = "
            f"{_money_br(second.regular_amortization, places)}",
        )
        start = first.reference_due_date + timedelta(days=1)
        _set_paragraph(
            paragraphs[67],
            f"O somatório dos juros diários, no período de {_date_br(start)} a "
            f"{_date_br(second.reference_due_date)}, totalizou {_money_br(second.interest_amount, places)}.",
        )
        _set_paragraph(
            paragraphs[71],
            f"P = {_money_br(second.regular_amortization, places)} + {_money_br(second.interest_amount, places)} "
            f"= {_money_br(second.installment_amount, places)}",
        )
        _set_paragraph(
            paragraphs[73],
            f"Dessa forma, a segunda prestação do contrato foi de {_money_br(second.installment_amount, places)}, "
            f"composta pela amortização de {_money_br(second.regular_amortization, places)} e pelos juros de "
            f"{_money_br(second.interest_amount, places)}.",
        )


def _fill_fixed_equations(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
) -> None:
    """Atualiza os valores das equações OMML sem destruir a diagramação original do Word."""
    first = result.installments[0]
    second = result.installments[1] if len(result.installments) > 1 else None
    paragraphs = doc.paragraphs
    places = settings.rounding.money_places
    pct_places = settings.rounding.percentage_display_places
    monthly_pct = _percent_br(result.monthly_interest_rate, pct_places)
    days = _period_day_count(result, 1)
    denominator = _single_competence_denominator(result, 1, settings)

    # Mantém as fórmulas originais do modelo, com notação matemática mais limpa.
    _set_math_texts(paragraphs[12], ["PMT= ", "VF × i", "1-(1+i)", "-n", " × ", "1", "1+i"])
    _set_math_texts(paragraphs[23], ["VF =", "Valor principal × ", "1+i", "n", " "])
    _set_math_texts(paragraphs[37], ["Juros da primeira prestação =", "1+i", "n", "-1", " × Valor principal"])
    _set_math_texts(paragraphs[47], ["Juros da segunda prestação = i × Valor principal"])

    # Mantém a fórmula genérica VF = Valor principal × (1+i)^n do modelo.
    if denominator is not None:
        _set_math_texts(
            paragraphs[24],
            [
                "VF =",
                f"{_money_br(first.opening_balance, places)} × ",
                "1+",
                monthly_pct,
                f"{days}/{denominator}",
                f"= {_money_br(first.balance_before_installment, places)}",
            ],
        )
    else:
        _set_paragraph(
            paragraphs[24],
            f"VF = {_money_br(first.opening_balance, places)} × fator acumulado dos juros diários "
            f"= {_money_br(first.balance_before_installment, places)}",
        )

    # Preserva a estrutura fracionária da fórmula PMT do modelo e troca apenas os valores.
    _set_math_texts(
        paragraphs[26],
        [
            "PMT= ",
            f"{_money_br(first.balance_before_installment, places)} × {monthly_pct}",
            f"1-(1+{monthly_pct})",
            f"-{settings.term}",
            " × ",
            "1",
            f"1+{monthly_pct}" if settings.payment_timing == 1 else "1",
        ],
    )
    _set_math_texts(
        paragraphs[28],
        [f"PMT = {_money_br(first.installment_amount, places)}"],
    )

    # Fórmula dos juros da primeira prestação: mantém o leiaute exponencial do modelo.
    if denominator is not None:
        _set_math_texts(
            paragraphs[38],
            [
                "Juros da primeira prestação =",
                f"1+{monthly_pct}",
                f"{days}/{denominator}",
                "-1",
                f"× {_money_br(first.opening_balance, places)} = {_money_br(first.interest_amount, places)}",
            ],
        )
    else:
        _set_paragraph(
            paragraphs[38],
            f"Juros da primeira prestação = soma dos juros diários do período = "
            f"{_money_br(first.interest_amount, places)}",
        )

    if second is not None:
        _set_math_texts(
            paragraphs[48],
            [
                f"Juros da segunda prestação = soma dos juros diários = "
                f"{_money_br(second.interest_amount, places)}"
            ],
        )


def _fill_fixed_narrative(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
) -> None:
    first = result.installments[0]
    second = result.installments[1] if len(result.installments) > 1 else None
    places = settings.rounding.money_places
    pct_places = settings.rounding.percentage_display_places
    paragraphs = doc.paragraphs
    days = _period_day_count(result, 1)
    start_interest = settings.credit_date + timedelta(days=1)

    _fill_fixed_equations(doc, settings, result)

    _set_paragraph(
        paragraphs[22],
        f"Primeiramente, apura-se o valor contratado na data da primeira prestação. Como a data do crédito "
        f"ocorreu em {_date_br(settings.credit_date)} e a primeira prestação foi calculada na data de "
        f"{_date_br(first.reference_due_date)}, foram considerados {days} dias de carência.",
    )
    _set_paragraph(
        paragraphs[29],
        f"Registra-se que a taxa de juros de {_percent_br(result.monthly_interest_rate, pct_places)} ao mês "
        f"é equivalente à taxa contratual de {_percent_br(settings.annual_interest_rate, pct_places)} ao ano.",
    )
    _set_paragraph(
        paragraphs[33],
        f"Portanto, para este contrato, os juros da primeira prestação correspondem aos {days} dias de carência.",
    )
    _set_paragraph(
        paragraphs[36],
        f"Assim, considerando a taxa anual de {_percent_br(settings.annual_interest_rate, pct_places)}, "
        f"equivalente a {_percent_br(result.monthly_interest_rate, pct_places)} ao mês, foram aplicados os "
        f"juros diários referentes ao período de {_date_br(start_interest)} a "
        f"{_date_br(first.reference_due_date)}.",
    )
    _set_paragraph(
        paragraphs[40],
        f"Dessa forma, a primeira prestação de {_money_br(first.installment_amount, places)} é composta por "
        f"juros de {_money_br(first.interest_amount, places)} e amortização de "
        f"{_money_br(first.regular_amortization, places)}, além do FGQC quando aplicável.",
    )
    _set_paragraph(
        paragraphs[44],
        f"O saldo principal de {_money_br(first.opening_balance, places)} foi reduzido exclusivamente pela "
        f"amortização de {_money_br(first.regular_amortization, places)}, que integra a primeira prestação. "
        f"Os juros de {_money_br(first.interest_amount, places)} remuneram o capital durante o período de "
        f"carência, mas não reduzem o principal. Após a amortização, o saldo devedor principal passou a "
        f"{_money_br(first.closing_balance, places)}.",
    )
    _set_paragraph(
        paragraphs[45],
        f"Para calcular os juros de cada parcela, utiliza-se a taxa mensal equivalente de "
        f"{_percent_br(result.monthly_interest_rate, pct_places)} (equivalente a "
        f"{_percent_br(settings.annual_interest_rate, pct_places)} ao ano).",
    )
    if second is not None:
        _set_paragraph(
            paragraphs[50],
            f"Na segunda prestação, os juros apurados foram de {_money_br(second.interest_amount, places)} e a "
            f"amortização foi de {_money_br(second.regular_amortization, places)}, totalizando a prestação "
            f"fixa de {_money_br(second.installment_amount, places)}. A redução do saldo principal decorre "
            f"da amortização, passando de {_money_br(second.opening_balance, places)} para "
            f"{_money_br(second.closing_balance, places)}.",
        )


def _hex_rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _set_cell_shading(cell: _Cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_margins(cell: _Cell, top: int = 80, start: int = 100, bottom: int = 80, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, margin_value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margin_value))
        node.set(qn("w:type"), "dxa")


def _set_cell_borders(cell: _Cell, color: str = MID_GRAY, size: str = "4") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = tc_borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _set_paragraph_shading(paragraph: Paragraph, fill: str) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    for old in p_pr.findall(qn("w:shd")):
        p_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def _set_paragraph_border(
    paragraph: Paragraph,
    edge: str,
    color: str,
    size: str = "12",
    space: str = "4",
) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.first_child_found_in("w:pBdr")
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    element = p_bdr.find(qn(f"w:{edge}"))
    if element is None:
        element = OxmlElement(f"w:{edge}")
        p_bdr.append(element)
    element.set(qn("w:val"), "single")
    element.set(qn("w:sz"), size)
    element.set(qn("w:space"), space)
    element.set(qn("w:color"), color)


def _set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _add_field(paragraph: Paragraph, instruction: str) -> None:
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = instruction
    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")
    fallback = OxmlElement("w:t")
    fallback.text = "1"
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.extend([fld_char_begin, instr_text, fld_char_separate, fallback, fld_char_end])


def _insert_title_banner(doc: DocumentObject, modality_name: str, contract_number: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.7)
    _set_cell_shading(cell, PETROLEUM)
    _set_cell_margins(cell, top=150, start=220, bottom=140, end=220)
    _set_cell_borders(cell, PETROLEUM, "0")

    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    title = p.add_run("PARECER TÉCNICO")
    title.bold = True
    title.font.name = "Arial"
    title.font.size = Pt(16)
    title.font.color.rgb = _hex_rgb(WHITE)

    subtitle = cell.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(0)
    r = subtitle.add_run(f"Evolução contratual  |  {modality_name}  |  Contrato {contract_number}")
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    r.font.color.rgb = _hex_rgb("D9ECF3")

    # Acento laranja discreto na base do banner.
    _set_paragraph_border(subtitle, "bottom", FUNCEF_ORANGE, size="20", space="5")

    table._tbl.getparent().remove(table._tbl)
    doc.paragraphs[0]._p.addprevious(table._tbl)


def _insert_summary_cards(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
    modality_name: str,
    anchor_table,
) -> None:
    first = result.installments[0]
    places = settings.rounding.money_places
    cards = [
        ("VALOR CONTRATADO", _money_br(settings.initial_balance, places)),
        ("VALOR NA 1ª PRESTAÇÃO", _money_br(first.balance_before_installment, places)),
        ("PRIMEIRA PRESTAÇÃO", _money_br(first.installment_amount, places)),
        ("JUROS DO PERÍODO", _money_br(first.interest_amount, places)),
        ("AMORTIZAÇÃO", _money_br(first.regular_amortization, places)),
        ("SALDO PRINCIPAL APÓS A 1ª", _money_br(first.closing_balance, places)),
    ]
    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_widths(table, [2.1, 2.1, 2.1])
    for idx, (label, value) in enumerate(cards):
        row, col = divmod(idx, 3)
        cell = table.cell(row, col)
        _set_cell_shading(cell, LIGHT_BLUE if row == 0 else LIGHTER_BLUE)
        _set_cell_borders(cell, WHITE, "8")
        _set_cell_margins(cell, top=110, start=130, bottom=100, end=130)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(3)
        label_run = p.add_run(label)
        label_run.bold = True
        label_run.font.name = "Arial"
        label_run.font.size = Pt(7.5)
        label_run.font.color.rgb = _hex_rgb(FUNCEF_BLUE)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(0)
        value_run = p2.add_run(value)
        value_run.bold = True
        value_run.font.name = "Arial"
        value_run.font.size = Pt(11)
        value_run.font.color.rgb = _hex_rgb(PETROLEUM)

    table._tbl.getparent().remove(table._tbl)
    anchor_table._tbl.addnext(table._tbl)


def _style_contract_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row_idx, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for col_idx, cell in enumerate(row.cells):
            _set_cell_margins(cell, top=75, start=100, bottom=75, end=100)
            _set_cell_borders(cell, MID_GRAY, "4")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_shading(cell, LIGHT_BLUE if col_idx == 0 else WHITE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    run.font.color.rgb = _hex_rgb(PETROLEUM if col_idx == 0 else DARK_TEXT)
                    run.bold = col_idx == 0


def _style_data_table(table, header_row: int = 0) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if table.rows:
        _set_repeat_table_header(table.rows[header_row])
    for row_idx, row in enumerate(table.rows):
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for cell in row.cells:
            _set_cell_margins(cell, top=55, start=55, bottom=55, end=55)
            _set_cell_borders(cell, MID_GRAY, "4")
            if row_idx == header_row:
                _set_cell_shading(cell, PETROLEUM)
            elif row_idx % 2 == 0:
                _set_cell_shading(cell, LIGHTER_BLUE)
            else:
                _set_cell_shading(cell, WHITE)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7.2)
                    run.font.color.rgb = _hex_rgb(WHITE if row_idx == header_row else DARK_TEXT)
                    run.bold = row_idx == header_row


def _style_rates_table(table) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for col_idx, cell in enumerate(row.cells):
            _set_cell_margins(cell, top=70, start=90, bottom=70, end=90)
            _set_cell_borders(cell, MID_GRAY, "4")
            _set_cell_shading(cell, LIGHT_BLUE if col_idx == 0 else WHITE)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_idx == 0 else WD_ALIGN_PARAGRAPH.RIGHT
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    run.bold = col_idx == 0
                    run.font.color.rgb = _hex_rgb(PETROLEUM if col_idx == 0 else DARK_TEXT)


def _style_formula_paragraph(paragraph: Paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.35)
    _set_paragraph_shading(paragraph, LIGHTER_BLUE)
    _set_paragraph_border(paragraph, "left", FUNCEF_BLUE, size="18", space="5")
    _set_paragraph_border(paragraph, "bottom", MID_GRAY, size="4", space="5")
    for run in paragraph.runs:
        run.font.name = "Cambria Math"
        run.font.size = Pt(10.5)
        run.font.color.rgb = _hex_rgb(PETROLEUM)


def _style_heading(paragraph: Paragraph) -> None:
    paragraph.paragraph_format.space_before = Pt(13)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    _set_paragraph_border(paragraph, "bottom", FUNCEF_ORANGE, size="10", space="3")
    for run in paragraph.runs:
        run.font.name = "Arial"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = _hex_rgb(PETROLEUM)


def _style_footer(doc: DocumentObject) -> None:
    for section in doc.sections:
        footer = section.footer
        if footer.tables:
            table = footer.tables[0]
        else:
            table = footer.add_table(rows=1, cols=2, width=Cm(16.7))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        _set_table_widths(table, [5.0, 1.3])
        left = table.cell(0, 0)
        right = table.cell(0, 1)
        _set_cell_borders(left, WHITE, "0")
        _set_cell_borders(right, WHITE, "0")
        lp = left.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        lp.paragraph_format.space_after = Pt(0)
        lr = lp.add_run("FUNCEF  •  Parecer técnico gerado pelo Motor de Cálculos")
        lr.font.name = "Arial"
        lr.font.size = Pt(7.5)
        lr.font.color.rgb = _hex_rgb("6B7A86")
        rp = right.paragraphs[0]
        rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        rp.paragraph_format.space_after = Pt(0)
        rr = rp.add_run("Página ")
        rr.font.name = "Arial"
        rr.font.size = Pt(7.5)
        rr.font.color.rgb = _hex_rgb("6B7A86")
        _add_field(rp, "PAGE")
        r2 = rp.add_run(" de ")
        r2.font.name = "Arial"
        r2.font.size = Pt(7.5)
        r2.font.color.rgb = _hex_rgb("6B7A86")
        _add_field(rp, "NUMPAGES")


def _apply_premium_style(
    doc: DocumentObject,
    settings: ContractSettings,
    result: EvolutionResult,
    modality_name: str,
    data: OpinionData,
) -> None:
    # Margens, tipografia e metadados.
    for section in doc.sections:
        section.top_margin = Cm(1.65)
        section.bottom_margin = Cm(1.55)
        section.left_margin = Cm(1.75)
        section.right_margin = Cm(1.75)
        section.header_distance = Cm(0.55)
        section.footer_distance = Cm(0.55)

    doc.core_properties.title = f"Parecer técnico - {modality_name} - Contrato {data.contract_number}"
    doc.core_properties.subject = "Evolução contratual e demonstração da metodologia de cálculo"
    doc.core_properties.keywords = "FUNCEF, parecer técnico, evolução contratual, cálculo"

    # Corpo textual: legibilidade corporativa, sem excesso de ornamentação.
    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.12
        if paragraph.paragraph_format.space_after is None:
            paragraph.paragraph_format.space_after = Pt(5)
        for run in paragraph.runs:
            if not ("m:oMath" in paragraph._p.xml or "m:oMathPara" in paragraph._p.xml):
                run.font.name = "Arial"
                run.font.size = Pt(10)
                run.font.color.rgb = _hex_rgb(DARK_TEXT)

    # Captura as tabelas originais antes de inserir os novos componentes visuais.
    contract_table = doc.tables[0]
    if settings.modality_code == "MOD_001":
        rates_table = None
        data_table = doc.tables[1]
    else:
        rates_table = doc.tables[1]
        data_table = doc.tables[2]

    _insert_title_banner(doc, modality_name, data.contract_number)
    _insert_summary_cards(doc, settings, result, modality_name, contract_table)

    # Tabelas: parâmetros, taxas e demonstrativos.
    _style_contract_table(contract_table)
    if rates_table is not None:
        _style_rates_table(rates_table)
    _style_data_table(data_table, 0)

    heading_prefixes = (
        "Ausência de Capitalização",
        "Aplicação do Sistema",
        "Cálculo da",
        "Cálculo dos",
        "Cálculo dos Juros",
        "Conclusão",
    )
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if any(text.startswith(prefix) for prefix in heading_prefixes):
            _style_heading(paragraph)
        is_math = "m:oMath" in paragraph._p.xml or "m:oMathPara" in paragraph._p.xml
        compact_formula = (
            text in {"A = PV / n", "J = SD × i", "P = A + J"}
            or (len(text) <= 125 and bool(re.match(r"^(A|J|P|VF|PMT)\s*=", text)) and "valor" not in text.lower())
        )
        if is_math or compact_formula:
            _style_formula_paragraph(paragraph)

    # Destaque da conclusão como bloco técnico.
    conclusion_found = False
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text == "Conclusão":
            conclusion_found = True
            continue
        if conclusion_found and text:
            _set_paragraph_shading(paragraph, LIGHTER_BLUE)
            _set_paragraph_border(paragraph, "left", FUNCEF_ORANGE, size="18", space="5")
            paragraph.paragraph_format.left_indent = Cm(0.25)
            paragraph.paragraph_format.right_indent = Cm(0.15)

    _style_footer(doc)


def build_opinion_docx(
    template_path: str | Path,
    settings: ContractSettings,
    result: EvolutionResult,
    modality_name: str,
    data: OpinionData,
) -> bytes:
    if len(result.installments) < 1:
        raise ValueError("A evolução não contém prestações para demonstrar no parecer.")
    doc = Document(template_path)
    intro = (
        f"O presente documento tem como objeto o contrato de empréstimo {modality_name} nº "
        f"{data.contract_number}, em nome de {data.participant_name}, matrícula {data.registration} "
        f"e CPF {data.cpf}, contratado nos seguintes parâmetros:"
    )
    _set_paragraph(doc.paragraphs[0], intro)
    _format_contract_table(doc, settings, result, modality_name, data)

    if settings.modality_code == "MOD_001":
        _fill_variable_summary_table(doc, settings, result, data)
        _fill_variable_narrative(doc, settings, result)
    elif settings.modality_code == "MOD_002":
        _fill_fixed_summary_tables(doc, settings, result, data)
        _fill_fixed_narrative(doc, settings, result)
    else:
        raise ValueError("Não há modelo de parecer configurado para esta modalidade.")

    _apply_premium_style(doc, settings, result, modality_name, data)

    output = BytesIO()
    doc.save(output)
    return output.getvalue()
