from __future__ import annotations

from dataclasses import dataclass, field, replace
from datetime import date, datetime, timedelta
import calendar
from decimal import Decimal, ROUND_HALF_UP, localcontext
from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .calendar_rules import (
    business_days_in_day_competence,
    calendar_days_in_day_competence,
    competence_bounds_for_day,
    competence_label_for_day,
    is_business_day,
    lagged_competence,
)
from .indices import IndexSeries
from .rates import monthly_to_daily_equivalent


PETROLEUM = "005A70"
ORANGE = "FF762B"
WHITE = "FFFFFF"
LIGHT = "EEF4F6"
THIN = Side(style="thin", color="D5E1E5")

CORRECTION_NONE = "SEM_CORRECAO"
CORRECTION_DAILY_BUSINESS = "DIARIA_UTEIS"
CORRECTION_DAILY_CALENDAR = "DIARIA_CORRIDOS"
CORRECTION_PRORATA_CALENDAR = "PRO_RATA_MENSAL_DIAS_CORRIDOS"
CORRECTION_MONTHLY_CLOSE = "MENSAL_FECHAMENTO"
CORRECTION_FACTOR_TABLE = "FATOR_TABELA"

INTEREST_NONE = "SEM_JUROS"
INTEREST_SIMPLE_30 = "SIMPLES_DIA_30"
INTEREST_SIMPLE_COMPETENCE = "SIMPLES_DIA_COMPETENCIA"
INTEREST_SIMPLE_360_TJRJ = "SIMPLES_360_TJRJ"
INTEREST_SIMPLE_YEARFRAC = "SIMPLES_FRACAOANO"
INTEREST_LEGAL_BCB = "TAXA_LEGAL_BCB"
INTEREST_COMPOUND_EQUIV = "COMPOSTO_DIARIO_EQUIVALENTE"

INTEREST_APPLICATION_DAILY = "APLICAR_DIARIAMENTE"
INTEREST_APPLICATION_END = "ACUMULAR_APLICAR_FINAL"

INTEREST_BASE_PRINCIPAL = "PRINCIPAL"
INTEREST_BASE_CORRECTED = "PRINCIPAL_CORRIGIDO"

ORDER_CORRECTION_INTEREST = "CORRECAO_JUROS"
ORDER_INTEREST_CORRECTION = "JUROS_CORRECAO"

PENALTY_BASE_PRINCIPAL = "PRINCIPAL"
PENALTY_BASE_CORRECTED = "PRINCIPAL_CORRIGIDO"
PENALTY_BASE_BALANCE = "SALDO_ANTES_MULTA"


def parse_ui_date(value) -> date:
    """Normaliza datas vindas da UI/CSV sem inverter dia e mês.

    Streamlit pode devolver DateColumn como ``date`` ou string ISO (YYYY-MM-DD).
    A aplicação também aceita o padrão brasileiro DD/MM/YYYY.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    if not text:
        raise ValueError("Data vazia.")
    try:
        return date.fromisoformat(text)
    except ValueError:
        pass
    for fmt in ("%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Data inválida: {text}. Use DD/MM/AAAA.")


@dataclass(frozen=True)
class UpdateItem:
    item_number: int
    origin_date: date
    amount: Decimal
    description: str = ""

    def __post_init__(self):
        if self.item_number <= 0:
            raise ValueError("O número do item/prestação deve ser positivo.")
        if self.amount <= 0:
            raise ValueError("O valor de cada item deve ser positivo.")


@dataclass(frozen=True)
class UpdateRule:
    order: int
    start_date: date
    end_date: date
    correction_index_code: str = CORRECTION_NONE
    correction_method: str = CORRECTION_NONE
    index_lag_months: int = 0
    monthly_interest_rate: Decimal = Decimal("0")
    interest_method: str = INTEREST_NONE
    interest_application: str = INTEREST_APPLICATION_DAILY
    interest_base: str = INTEREST_BASE_CORRECTED
    event_order: str = ORDER_CORRECTION_INTEREST
    competence_start_day: int = 1
    interest_start_date: date | None = None
    yearfrac_basis: int = 0
    description: str = ""

    def __post_init__(self):
        if self.order <= 0:
            raise ValueError("A ordem da regra deve ser positiva.")
        if self.end_date < self.start_date:
            raise ValueError(f"Regra {self.order}: a data final não pode ser anterior à inicial.")
        if self.index_lag_months < 0:
            raise ValueError(f"Regra {self.order}: a defasagem não pode ser negativa.")
        if self.monthly_interest_rate <= Decimal("-1"):
            raise ValueError(f"Regra {self.order}: a taxa mensal deve ser superior a -100%.")
        if self.competence_start_day not in (1, 21):
            raise ValueError(f"Regra {self.order}: o início da competência deve ser 1 ou 21.")
        if self.correction_method not in {
            CORRECTION_NONE, CORRECTION_DAILY_BUSINESS,
            CORRECTION_DAILY_CALENDAR, CORRECTION_PRORATA_CALENDAR, CORRECTION_MONTHLY_CLOSE, CORRECTION_FACTOR_TABLE,
        }:
            raise ValueError(f"Regra {self.order}: método de correção inválido.")
        if self.interest_method not in {
            INTEREST_NONE, INTEREST_SIMPLE_30,
            INTEREST_SIMPLE_COMPETENCE, INTEREST_SIMPLE_360_TJRJ, INTEREST_SIMPLE_YEARFRAC, INTEREST_LEGAL_BCB, INTEREST_COMPOUND_EQUIV,
        }:
            raise ValueError(f"Regra {self.order}: método de juros inválido.")
        if self.interest_application not in {INTEREST_APPLICATION_DAILY, INTEREST_APPLICATION_END}:
            raise ValueError(f"Regra {self.order}: forma de aplicação dos juros inválida.")
        if self.interest_method == INTEREST_COMPOUND_EQUIV and self.interest_application != INTEREST_APPLICATION_DAILY:
            raise ValueError(f"Regra {self.order}: juros compostos equivalentes exigem aplicação diária.")
        if self.yearfrac_basis not in {0, 1, 2, 3, 4}:
            raise ValueError(f"Regra {self.order}: base FRAÇÃOANO deve estar entre 0 e 4.")
        if self.interest_base not in {INTEREST_BASE_PRINCIPAL, INTEREST_BASE_CORRECTED}:
            raise ValueError(f"Regra {self.order}: base de juros inválida.")
        if self.event_order not in {ORDER_CORRECTION_INTEREST, ORDER_INTEREST_CORRECTION}:
            raise ValueError(f"Regra {self.order}: ordem de incidência inválida.")


@dataclass(frozen=True)
class Penalty:
    event_date: date
    rate: Decimal
    base: str = PENALTY_BASE_BALANCE
    target_item: int = 0
    note: str = ""

    def __post_init__(self):
        if self.rate <= Decimal("-1"):
            raise ValueError("A multa deve ser superior a -100%.")
        if self.rate < 0:
            raise ValueError("A multa não pode ser negativa.")
        if self.target_item < 0:
            raise ValueError("O item alvo da multa deve ser zero (todos) ou positivo.")
        if self.base not in {PENALTY_BASE_PRINCIPAL, PENALTY_BASE_CORRECTED, PENALTY_BASE_BALANCE}:
            raise ValueError("Base de multa inválida.")


@dataclass(frozen=True)
class Abatement:
    event_date: date
    amount: Decimal
    target_item: int = 0  # 0 = todas as obrigações abertas
    note: str = ""

    def __post_init__(self):
        if self.amount <= 0:
            raise ValueError("O abatimento deve ser positivo.")
        if self.target_item < 0:
            raise ValueError("O item alvo deve ser zero (todas) ou um número positivo.")


@dataclass
class _ItemState:
    item: UpdateItem
    principal: Decimal
    correction: Decimal = Decimal("0")
    interest: Decimal = Decimal("0")
    penalty: Decimal = Decimal("0")
    total_abatements: Decimal = Decimal("0")
    total_correction_generated: Decimal = Decimal("0")
    total_interest_generated: Decimal = Decimal("0")
    total_penalty_generated: Decimal = Decimal("0")

    @property
    def corrected_principal(self) -> Decimal:
        return self.principal + self.correction

    @property
    def total(self) -> Decimal:
        return self.principal + self.correction + self.interest + self.penalty


@dataclass(frozen=True)
class UpdateDailyRow:
    day: date
    item_number: int
    rule_order: int | None
    rule_description: str
    competence: str | None
    index_reference: str | None
    is_business_day: bool
    opening_principal: Decimal
    opening_correction: Decimal
    opening_interest: Decimal
    opening_penalty: Decimal
    correction_rate: Decimal
    correction_amount: Decimal
    interest_rate: Decimal
    interest_amount: Decimal
    penalty_rate: Decimal
    penalty_amount: Decimal
    abatement_amount: Decimal
    abatement_principal: Decimal
    abatement_correction: Decimal
    abatement_interest: Decimal
    abatement_penalty: Decimal
    closing_principal: Decimal
    closing_correction: Decimal
    closing_interest: Decimal
    closing_penalty: Decimal
    closing_total: Decimal


@dataclass(frozen=True)
class UpdateItemSummary:
    item_number: int
    description: str
    origin_date: date
    original_amount: Decimal
    correction_amount: Decimal
    interest_amount: Decimal
    penalty_amount: Decimal
    abatements: Decimal
    updated_total: Decimal


@dataclass
class MonetaryUpdateResult:
    base_date: date
    items: list[UpdateItemSummary]
    daily_rows: list[UpdateDailyRow]
    rules: list[UpdateRule]
    abatements: list[Abatement]
    penalties: list[Penalty] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    is_complete: bool = True
    stopped_at: date | None = None
    missing_index_code: str | None = None
    missing_index_reference: str | None = None

    @property
    def original_total(self) -> Decimal:
        return sum((x.original_amount for x in self.items), Decimal("0"))

    @property
    def correction_total(self) -> Decimal:
        return sum((x.correction_amount for x in self.items), Decimal("0"))

    @property
    def interest_total(self) -> Decimal:
        return sum((x.interest_amount for x in self.items), Decimal("0"))

    @property
    def penalty_total(self) -> Decimal:
        return sum((x.penalty_amount for x in self.items), Decimal("0"))

    @property
    def abatements_total(self) -> Decimal:
        return sum((x.abatements for x in self.items), Decimal("0"))

    @property
    def updated_total(self) -> Decimal:
        return sum((x.updated_total for x in self.items), Decimal("0"))

    @property
    def correction_through_date(self) -> date:
        """Última data alcançada pela correção monetária sem projetar índice futuro."""
        if self.is_complete or self.stopped_at is None:
            return self.base_date
        return self.stopped_at - timedelta(days=1)


class MissingIndexError(KeyError):
    def __init__(self, day: date, code: str, reference: str):
        self.day = day
        self.code = code
        self.reference = reference
        super().__init__(f"Índice {code} ausente para a competência {reference}.")


def _quantize_money(value: Decimal, places: int) -> Decimal:
    quantum = Decimal(1).scaleb(-places)
    return value.quantize(quantum, rounding=ROUND_HALF_UP)


def _equivalent_daily(monthly_rate: Decimal, days: int) -> Decimal:
    if days <= 0:
        return Decimal("0")
    with localcontext() as ctx:
        ctx.prec = 40
        return (Decimal("1") + monthly_rate) ** (Decimal("1") / Decimal(days)) - Decimal("1")


def tjrj_commercial_days_30_360(start_date: date, end_date: date) -> int:
    """Contagem comercial usada pelo SCJ WEB/TJRJ para juros simples.

    O manual oficial confirma o ano comercial de 360 dias. A calculadora observada
    também mostra dois comportamentos de borda que precisam ser reproduzidos:
    06/06/2023 a 06/06/2025 = 720 dias; 02/05/2018 a 07/08/2026 = 2.975 dias;
    e, quando o termo final é o dia 30, 02/05/2018 a 30/06/2026 = 2.939 dias.
    Assim, parte-se da convenção 30/360 e inclui-se o termo final quando ele cai
    no dia comercial 30.
    """
    if end_date < start_date:
        return 0
    d1 = min(start_date.day, 30)
    d2 = min(end_date.day, 30)
    days = (end_date.year - start_date.year) * 360
    days += (end_date.month - start_date.month) * 30
    days += d2 - d1
    if end_date.day >= 30:
        days += 1
    return max(0, days)


def tjrj_simple_360_rate(monthly_rate: Decimal, start_date: date, end_date: date) -> Decimal:
    """Taxa simples acumulada: taxa mensal x dias comerciais / 30."""
    days = tjrj_commercial_days_30_360(start_date, end_date)
    return monthly_rate * Decimal(days) / Decimal("30")


def _find_rule(day: date, rules: list[UpdateRule]) -> UpdateRule | None:
    candidates = [rule for rule in rules if rule.start_date <= day <= rule.end_date]
    if not candidates:
        return None
    candidates.sort(key=lambda x: (x.order, x.start_date))
    return candidates[-1]


def validate_rules(rules: Iterable[UpdateRule]) -> list[UpdateRule]:
    ordered = sorted(rules, key=lambda x: (x.start_date, x.order, x.end_date))
    if not ordered:
        raise ValueError("Cadastre ao menos uma regra de atualização.")
    # Sobreposição é aceita apenas quando a ordem resolve a prioridade; registramos isso
    # de forma determinística em _find_rule. Datas totalmente fora da apuração são inofensivas.
    return ordered


def _correction_rate_for_day(
    day: date,
    rule: UpdateRule,
    indices: dict[str, IndexSeries],
    holidays: set[date],
    factor_series: dict[str, object] | None = None,
) -> tuple[Decimal, str | None, str]:
    competence = competence_label_for_day(day, rule.competence_start_day)
    if rule.correction_method == CORRECTION_NONE or rule.correction_index_code in {"", CORRECTION_NONE, None}:
        return Decimal("0"), None, competence

    code = rule.correction_index_code

    if rule.correction_method == CORRECTION_FACTOR_TABLE:
        # Tabelas práticas judiciais são tabelas de fatores mensais. A incidência é
        # registrada no primeiro dia de cada novo mês, de modo que o produto das
        # razões mensais reproduza exatamente F(mês final) / F(mês inicial).
        current_ref = f"{day.year:04d}-{day.month:02d}"
        if day.day != 1:
            return Decimal("0"), current_ref, competence
        previous_day = day - timedelta(days=1)
        previous_ref = f"{previous_day.year:04d}-{previous_day.month:02d}"
        series_map = factor_series or {}
        if code not in series_map:
            raise MissingIndexError(day, code, current_ref)
        values = getattr(series_map[code], "values", {})

        # O SCJ WEB/TJRJ aceita a data atual mesmo quando o mês corrente ainda
        # não possui novo fator mensal. Nesse caso, o cálculo permanece com o
        # último fator mensal disponível. Ex.: em 07/08/2026, o fator vigente
        # é o de 07/2026. O fallback é restrito ao mês imediatamente posterior
        # à última competência disponível, evitando carregar fator antigo por
        # vários meses sem atualização.
        if code == "TJRJ_CIVEL_14905" and current_ref not in values:
            refs = sorted(values)
            latest_ref = refs[-1] if refs else None
            if latest_ref:
                ly, lm = map(int, latest_ref.split("-"))
                ny, nm = (ly + 1, 1) if lm == 12 else (ly, lm + 1)
                if current_ref == f"{ny:04d}-{nm:02d}" and previous_ref == latest_ref:
                    return Decimal("0"), latest_ref, competence

        if previous_ref not in values:
            raise MissingIndexError(day, code, previous_ref)
        if current_ref not in values:
            raise MissingIndexError(day, code, current_ref)
        previous_factor = Decimal(str(values[previous_ref]))
        current_factor = Decimal(str(values[current_ref]))
        if previous_factor == 0:
            raise ValueError(f"Tabela {code}: fator zero na competência {previous_ref}.")
        return current_factor / previous_factor - Decimal("1"), current_ref, competence

    # Critério livre de pró-rata mensal: cada cotação mensal é transformada em
    # fator diário equivalente pelo número de dias CORRIDOS do próprio mês.
    # Assim, em uma fração de mês com d dias de um mês de D dias, o fator efetivo
    # é (1 + i_mensal) ** (d / D). A data de origem é incluída no período,
    # reproduzindo o critério da planilha operacional utilizada como referência.
    if rule.correction_method == CORRECTION_PRORATA_CALENDAR:
        competence = f"{day.year:04d}-{day.month:02d}"
        reference = lagged_competence(competence, rule.index_lag_months)
        if code not in indices or reference not in indices[code].values:
            raise MissingIndexError(day, code, reference)
        monthly_rate = indices[code].values[reference]
        days = calendar.monthrange(day.year, day.month)[1]
        return _equivalent_daily(monthly_rate, days), reference, competence

    reference = lagged_competence(competence, rule.index_lag_months)
    if code not in indices or reference not in indices[code].values:
        raise MissingIndexError(day, code, reference)
    monthly_rate = indices[code].values[reference]

    if rule.correction_method == CORRECTION_DAILY_BUSINESS:
        if not is_business_day(day, holidays):
            return Decimal("0"), reference, competence
        days = business_days_in_day_competence(day, rule.competence_start_day, holidays)
        return _equivalent_daily(monthly_rate, days), reference, competence

    if rule.correction_method == CORRECTION_DAILY_CALENDAR:
        days = calendar_days_in_day_competence(day, rule.competence_start_day)
        return _equivalent_daily(monthly_rate, days), reference, competence

    if rule.correction_method == CORRECTION_MONTHLY_CLOSE:
        _, end = competence_bounds_for_day(day, rule.competence_start_day)
        return (monthly_rate if day == end else Decimal("0")), reference, competence

    return Decimal("0"), reference, competence


def _is_last_day_of_february(value: date) -> bool:
    return value.month == 2 and value.day == calendar.monthrange(value.year, value.month)[1]


def _days360_us_nasd(start: date, end: date) -> int:
    """Dias 30/360 no padrão US/NASD usado pela FRAÇÃOANO base 0 do Excel."""
    y1, m1, d1 = start.year, start.month, start.day
    y2, m2, d2 = end.year, end.month, end.day

    if d1 == 31 and d2 == 31:
        d1 = d2 = 30
    elif d1 == 31:
        d1 = 30
    elif d1 == 30 and d2 == 31:
        d2 = 30
    elif _is_last_day_of_february(start) and _is_last_day_of_february(end):
        d1 = d2 = 30
    elif _is_last_day_of_february(start):
        d1 = 30

    return (y2 - y1) * 360 + (m2 - m1) * 30 + (d2 - d1)


def _days360_european(start: date, end: date) -> int:
    d1 = min(start.day, 30)
    d2 = min(end.day, 30)
    return (end.year - start.year) * 360 + (end.month - start.month) * 30 + (d2 - d1)


def _is_leap(year: int) -> bool:
    return calendar.isleap(year)


def _feb29_between(start: date, end: date) -> bool:
    for year in range(start.year, end.year + 1):
        if not _is_leap(year):
            continue
        leap_day = date(year, 2, 29)
        if start <= leap_day < end:
            return True
    return False


def excel_yearfrac(start: date, end: date, basis: int = 0) -> Decimal:
    """Replica a FRAÇÃOANO/YEARFRAC do Excel para as bases 0 a 4.

    A função é usada para juros simples: o valor retornado representa a fração
    de ano entre as datas, sem capitalização.
    """
    if basis not in {0, 1, 2, 3, 4}:
        raise ValueError("Base FRAÇÃOANO inválida. Use 0, 1, 2, 3 ou 4.")
    if start == end:
        return Decimal("0")
    sign = Decimal("1")
    if start > end:
        start, end = end, start
        sign = Decimal("-1")

    if basis == 0:
        return sign * Decimal(_days360_us_nasd(start, end)) / Decimal("360")
    if basis == 2:
        return sign * Decimal((end - start).days) / Decimal("360")
    if basis == 3:
        return sign * Decimal((end - start).days) / Decimal("365")
    if basis == 4:
        return sign * Decimal(_days360_european(start, end)) / Decimal("360")

    # Base 1 - Real/Real, seguindo o comportamento documentado do Excel.
    same_or_within_one_year = (
        start.year == end.year
        or (end.year == start.year + 1 and (start.month > end.month or (start.month == end.month and start.day >= end.day)))
    )
    if same_or_within_one_year:
        year_length = Decimal("365")
        if (start.year == end.year and _is_leap(start.year)) or _feb29_between(start, end) or (end.month == 2 and end.day == 29):
            year_length = Decimal("366")
        return sign * Decimal((end - start).days) / year_length

    years = end.year - start.year + 1
    total_days = (date(end.year + 1, 1, 1) - date(start.year, 1, 1)).days
    average_year = Decimal(total_days) / Decimal(years)
    return sign * Decimal((end - start).days) / average_year


def yearfrac_simple_rate(monthly_rate: Decimal, start: date, end: date, basis: int = 0) -> Decimal:
    """Taxa simples acumulada = taxa anual (taxa mensal x 12) x FRAÇÃOANO."""
    if end < start or monthly_rate == 0:
        return Decimal("0")
    return monthly_rate * Decimal("12") * excel_yearfrac(start, end, basis)


def legal_rate_simple_accumulated(series: IndexSeries, start: date, end: date) -> Decimal:
    """Acumula a Taxa Legal (BCB/CMN 5.171) em regime simples.

    O intervalo segue a lógica da Calculadora do Cidadão: data inicial inclusiva e
    data final exclusiva. Meses completos somam a TL mensal; frações de mês usam
    pro rata por dias corridos, com o percentual proporcional arredondado a 6
    casas decimais antes da conversão para taxa decimal.
    """
    start = max(start, date(2024, 8, 30))
    if end <= start:
        return Decimal("0")
    total = Decimal("0")
    cursor = start
    while cursor < end:
        key = f"{cursor.year:04d}-{cursor.month:02d}"
        if key not in series.values:
            raise MissingIndexError(cursor, series.code, key)
        month_end_exclusive = (date(cursor.year + 1, 1, 1) if cursor.month == 12 else date(cursor.year, cursor.month + 1, 1))
        segment_end = min(end, month_end_exclusive)
        days = (segment_end - cursor).days
        days_in_month = calendar.monthrange(cursor.year, cursor.month)[1]
        monthly_rate = series.values[key]
        # Res. CMN 5.171: juros simples + fração pro rata por dias corridos.
        proportional_pct = (monthly_rate * Decimal("100") * Decimal(days) / Decimal(days_in_month)).quantize(
            Decimal("0.000001"), rounding=ROUND_HALF_UP
        )
        total += proportional_pct / Decimal("100")
        cursor = segment_end
    return total


def _interest_rate_for_day(day: date, rule: UpdateRule) -> Decimal:
    effective_start = rule.interest_start_date or rule.start_date
    if day < effective_start:
        return Decimal("0")
    monthly = rule.monthly_interest_rate
    if rule.interest_method == INTEREST_NONE or monthly == 0:
        return Decimal("0")
    if rule.interest_method == INTEREST_SIMPLE_30:
        return monthly / Decimal("30")
    if rule.interest_method == INTEREST_SIMPLE_COMPETENCE:
        days = calendar_days_in_day_competence(day, rule.competence_start_day)
        return monthly / Decimal(days)
    if rule.interest_method == INTEREST_SIMPLE_360_TJRJ:
        # O perfil TJRJ calcula o percentual total no fechamento da regra.
        return Decimal("0")
    if rule.interest_method == INTEREST_LEGAL_BCB:
        # Taxa Legal é acumulada em regime simples no fechamento.
        return Decimal("0")
    if rule.interest_method == INTEREST_SIMPLE_YEARFRAC:
        # FRAÇÃOANO também é um método de fechamento: aplica a taxa simples
        # acumulada ao saldo/base existente na data final da regra.
        return Decimal("0")
    if rule.interest_method == INTEREST_COMPOUND_EQUIV:
        days = calendar_days_in_day_competence(day, rule.competence_start_day)
        return monthly_to_daily_equivalent(monthly, days)
    return Decimal("0")


def _apply_correction(state: _ItemState, rate: Decimal, round_daily: bool, money_places: int) -> Decimal:
    if rate == 0 or state.corrected_principal <= 0:
        return Decimal("0")
    amount = state.corrected_principal * rate
    if round_daily:
        amount = _quantize_money(amount, money_places)
    state.correction += amount
    state.total_correction_generated += amount
    if round_daily:
        state.correction = _quantize_money(state.correction, money_places)
    return amount


def _apply_interest(state: _ItemState, rate: Decimal, rule: UpdateRule, round_daily: bool, money_places: int) -> Decimal:
    if rate == 0 or state.total <= 0:
        return Decimal("0")
    base = state.principal if rule.interest_base == INTEREST_BASE_PRINCIPAL else state.corrected_principal
    # No método composto, os juros acumulados integram a base do dia seguinte.
    # Nos métodos simples, os juros anteriores ficam fora da base.
    if rule.interest_method == INTEREST_COMPOUND_EQUIV:
        base += state.interest
    if base <= 0:
        return Decimal("0")
    amount = base * rate
    if round_daily:
        amount = _quantize_money(amount, money_places)
    state.interest += amount
    state.total_interest_generated += amount
    if round_daily:
        state.interest = _quantize_money(state.interest, money_places)
    return amount


def _apply_penalty(state: _ItemState, penalty: Penalty, round_daily: bool, money_places: int) -> Decimal:
    if penalty.rate == 0 or state.total <= 0:
        return Decimal("0")
    if penalty.base == PENALTY_BASE_PRINCIPAL:
        base = state.principal
    elif penalty.base == PENALTY_BASE_CORRECTED:
        base = state.corrected_principal
    else:
        base = state.principal + state.correction + state.interest
    if base <= 0:
        return Decimal("0")
    amount = base * penalty.rate
    if round_daily:
        amount = _quantize_money(amount, money_places)
    state.penalty += amount
    state.total_penalty_generated += amount
    if round_daily:
        state.penalty = _quantize_money(state.penalty, money_places)
    return amount


def _reduce_proportionally(state: _ItemState, amount: Decimal, money_places: int) -> tuple[Decimal, Decimal, Decimal, Decimal, Decimal]:
    total = state.total
    if total <= 0 or amount <= 0:
        return Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0")
    applied = min(amount, total)
    ratio = applied / total

    p = _quantize_money(state.principal * ratio, money_places)
    c = _quantize_money(state.correction * ratio, money_places)
    i = _quantize_money(state.interest * ratio, money_places)
    m = _quantize_money(state.penalty * ratio, money_places)
    distributed = p + c + i + m
    residue = applied - distributed

    # Ajusta resíduo de arredondamento no maior componente ainda disponível.
    components = [
        ("principal", state.principal - p),
        ("correction", state.correction - c),
        ("interest", state.interest - i),
        ("penalty", state.penalty - m),
    ]
    components.sort(key=lambda x: x[1], reverse=True)
    if residue != 0:
        name = components[0][0]
        if name == "principal":
            p += residue
        elif name == "correction":
            c += residue
        elif name == "interest":
            i += residue
        else:
            m += residue

    state.principal = max(Decimal("0"), state.principal - p)
    state.correction = max(Decimal("0"), state.correction - c)
    state.interest = max(Decimal("0"), state.interest - i)
    state.penalty = max(Decimal("0"), state.penalty - m)
    state.total_abatements += applied
    return applied, p, c, i, m


def _apply_abatement_to_states(
    states: list[_ItemState],
    abatement: Abatement,
    money_places: int,
) -> tuple[dict[int, tuple[Decimal, Decimal, Decimal, Decimal, Decimal]], Decimal]:
    eligible = [
        state for state in states
        if state.item.origin_date < abatement.event_date and state.total > 0
        and (abatement.target_item == 0 or state.item.item_number == abatement.target_item)
    ]
    allocations: dict[int, tuple[Decimal, Decimal, Decimal, Decimal, Decimal]] = {}
    if not eligible:
        return allocations, abatement.amount

    total_open = sum((state.total for state in eligible), Decimal("0"))
    amount_to_apply = min(abatement.amount, total_open)
    remaining = amount_to_apply

    for idx, state in enumerate(eligible):
        if idx == len(eligible) - 1:
            share = remaining
        else:
            share = _quantize_money(amount_to_apply * state.total / total_open, money_places)
            share = min(share, remaining)
        allocations[state.item.item_number] = _reduce_proportionally(state, share, money_places)
        remaining -= allocations[state.item.item_number][0]

    not_applied = abatement.amount - amount_to_apply
    return allocations, not_applied


def calculate_monetary_update(
    items: Iterable[UpdateItem],
    rules: Iterable[UpdateRule],
    base_date: date,
    indices: dict[str, IndexSeries],
    holidays: set[date],
    abatements: Iterable[Abatement] = (),
    penalties: Iterable[Penalty] = (),
    factor_series: dict[str, object] | None = None,
    *,
    round_daily_money: bool = True,
    money_places: int = 2,
) -> MonetaryUpdateResult:
    items_list = sorted(items, key=lambda x: (x.origin_date, x.item_number))
    if not items_list:
        raise ValueError("Informe ao menos um valor para atualização.")
    if base_date <= min(x.origin_date for x in items_list):
        raise ValueError("A data-base precisa ser posterior à data de origem de pelo menos um valor.")
    rules_list = validate_rules(rules)
    abatements_list = sorted(abatements, key=lambda x: (x.event_date, x.target_item))
    penalties_list = sorted(penalties, key=lambda x: (x.event_date, x.target_item))
    if any(x.event_date > base_date for x in abatements_list):
        raise ValueError("Existem abatimentos posteriores à data-base.")
    if any(x.event_date > base_date for x in penalties_list):
        raise ValueError("Existem multas posteriores à data-base.")

    states = [_ItemState(item=x, principal=x.amount) for x in items_list]
    daily_rows: list[UpdateDailyRow] = []
    warnings: list[str] = []
    missing_gap_dates: set[date] = set()
    abatements_by_date: dict[date, list[Abatement]] = {}
    for ab in abatements_list:
        abatements_by_date.setdefault(ab.event_date, []).append(ab)
    penalties_by_date: dict[date, list[Penalty]] = {}
    for pen in penalties_list:
        penalties_by_date.setdefault(pen.event_date, []).append(pen)

    start = min(x.origin_date for x in items_list)
    is_complete = True
    stopped_at = None
    missing_code = None
    missing_ref = None
    deferred_interest_rates: dict[tuple[int, int], Decimal] = {}
    # Quando uma série de correção deixa de ter competência disponível, a correção
    # daquela regra é congelada a partir do primeiro mês ausente. Juros e demais
    # eventos continuam até a data-base; nenhum índice futuro é projetado.
    frozen_correction_rules: dict[int, tuple[date, str, str]] = {}

    day = start
    while day <= base_date:
        rule = _find_rule(day, rules_list)
        include_origin_day = bool(rule is not None and rule.correction_method == CORRECTION_PRORATA_CALENDAR)
        active_states = [
            state for state in states
            if state.total > 0 and (state.item.origin_date < day or (include_origin_day and state.item.origin_date == day))
        ]
        if not active_states:
            day += timedelta(days=1)
            continue

        if rule is None:
            missing_gap_dates.add(day)
            correction_rate = Decimal("0")
            interest_rate = Decimal("0")
            index_reference = None
            competence = None
        else:
            if rule.order in frozen_correction_rules:
                _, frozen_code, frozen_ref = frozen_correction_rules[rule.order]
                correction_rate = Decimal("0")
                index_reference = frozen_ref
                competence = competence_label_for_day(day, rule.competence_start_day)
            else:
                try:
                    correction_rate, index_reference, competence = _correction_rate_for_day(
                        day, rule, indices, holidays, factor_series
                    )
                except MissingIndexError as exc:
                    is_complete = False
                    if stopped_at is None or exc.day < stopped_at:
                        stopped_at = exc.day
                        missing_code = exc.code
                        missing_ref = exc.reference
                    frozen_correction_rules[rule.order] = (exc.day, exc.code, exc.reference)
                    correction_rate = Decimal("0")
                    index_reference = exc.reference
                    competence = competence_label_for_day(day, rule.competence_start_day)
                    warnings.append(
                        f"Correção monetária limitada em {exc.day.strftime('%d/%m/%Y')}: índice {exc.code} "
                        f"não disponível para a competência {exc.reference}. A correção não foi projetada; "
                        "o saldo corrigido permaneceu congelado e os juros configurados continuaram até a data-base."
                    )
            interest_rate = _interest_rate_for_day(day, rule)

        day_rows: dict[int, dict] = {}
        for state in active_states:
            opening = {
                "principal": state.principal,
                "correction": state.correction,
                "interest": state.interest,
                "penalty": state.penalty,
            }
            correction_amount = Decimal("0")
            interest_amount = Decimal("0")
            state_interest_rate = interest_rate
            if day <= state.item.origin_date:
                state_interest_rate = Decimal("0")
            correction_rounding = bool(
                round_daily_money and rule is not None and rule.correction_method not in {CORRECTION_FACTOR_TABLE, CORRECTION_PRORATA_CALENDAR}
            )

            # Há duas formas de aplicação para juros simples:
            # 1) aplicar monetariamente a cada dia sobre a base daquele dia;
            # 2) apenas acumular o percentual do período e aplicar uma única vez
            #    no fechamento da regra, sobre a base já corrigida.
            # TJRJ/360 e FRAÇÃOANO são, por definição, métodos de fechamento.
            apply_at_end = bool(
                rule is not None
                and rule.interest_method != INTEREST_NONE
                and (
                    rule.interest_application == INTEREST_APPLICATION_END
                    or rule.interest_method in {INTEREST_SIMPLE_360_TJRJ, INTEREST_SIMPLE_YEARFRAC, INTEREST_LEGAL_BCB}
                )
            )
            if rule is not None and apply_at_end:
                terminal_day = min(rule.end_date, base_date)
                if rule.interest_method in {INTEREST_SIMPLE_360_TJRJ, INTEREST_SIMPLE_YEARFRAC, INTEREST_LEGAL_BCB}:
                    effective_start = max(rule.interest_start_date or rule.start_date, state.item.origin_date)
                    if day == terminal_day and terminal_day >= effective_start and (rule.monthly_interest_rate != 0 or rule.interest_method == INTEREST_LEGAL_BCB):
                        if rule.interest_method == INTEREST_SIMPLE_360_TJRJ:
                            state_interest_rate = tjrj_simple_360_rate(rule.monthly_interest_rate, effective_start, terminal_day)
                        elif rule.interest_method == INTEREST_SIMPLE_YEARFRAC:
                            state_interest_rate = yearfrac_simple_rate(rule.monthly_interest_rate, effective_start, terminal_day, rule.yearfrac_basis)
                        else:
                            legal_series = indices.get("TAXA_LEGAL")
                            if legal_series is None:
                                raise MissingIndexError(terminal_day, "TAXA_LEGAL", f"{terminal_day.year:04d}-{terminal_day.month:02d}")
                            state_interest_rate = legal_rate_simple_accumulated(legal_series, effective_start, terminal_day)
                    else:
                        state_interest_rate = Decimal("0")
                else:
                    # Nos métodos simples diários, acumulamos apenas a TAXA diária.
                    # O valor monetário dos juros é gerado uma única vez no final.
                    first_interest_day = max(
                        rule.interest_start_date or rule.start_date,
                        state.item.origin_date + timedelta(days=1),
                    )
                    key = (state.item.item_number, rule.order)
                    if day >= first_interest_day and interest_rate != 0:
                        deferred_interest_rates[key] = deferred_interest_rates.get(key, Decimal("0")) + interest_rate
                    state_interest_rate = deferred_interest_rates.get(key, Decimal("0")) if day == terminal_day else Decimal("0")

            # Quando os juros são aplicados no fechamento, a correção deve estar
            # completamente apurada antes da incidência sobre o saldo corrigido.
            # A ordem configurável continua valendo somente para aplicação diária.
            if rule is not None and not apply_at_end and rule.event_order == ORDER_INTEREST_CORRECTION:
                interest_amount = _apply_interest(state, state_interest_rate, rule, round_daily_money, money_places)
                correction_amount = _apply_correction(state, correction_rate, correction_rounding, money_places)
            else:
                if rule is not None:
                    correction_amount = _apply_correction(state, correction_rate, correction_rounding, money_places)
                    interest_amount = _apply_interest(state, state_interest_rate, rule, round_daily_money, money_places)

            day_rows[state.item.item_number] = {
                "opening": opening,
                "correction_amount": correction_amount,
                "interest_rate": state_interest_rate,
                "interest_amount": interest_amount,
                "penalty_rate": Decimal("0"),
                "penalty_amount": Decimal("0"),
                "ab": (Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0"), Decimal("0")),
            }

        for pen in penalties_by_date.get(day, []):
            eligible = [
                state for state in states
                if state.item.origin_date < day and state.total > 0
                and (pen.target_item == 0 or state.item.item_number == pen.target_item)
            ]
            for state in eligible:
                amount = _apply_penalty(state, pen, round_daily_money, money_places)
                if state.item.item_number in day_rows:
                    day_rows[state.item.item_number]["penalty_rate"] += pen.rate
                    day_rows[state.item.item_number]["penalty_amount"] += amount

        for ab in abatements_by_date.get(day, []):
            allocations, not_applied = _apply_abatement_to_states(states, ab, money_places)
            if not_applied > 0:
                warnings.append(
                    f"Abatimento de {ab.event_date.strftime('%d/%m/%Y')} teve excedente não aplicado de "
                    f"R$ {not_applied:.2f}."
                )
            for item_number, allocation in allocations.items():
                if item_number in day_rows:
                    previous = day_rows[item_number]["ab"]
                    day_rows[item_number]["ab"] = tuple(previous[i] + allocation[i] for i in range(5))

        for state in active_states:
            data = day_rows[state.item.item_number]
            applied, ab_p, ab_c, ab_i, ab_m = data["ab"]
            daily_rows.append(UpdateDailyRow(
                day=day,
                item_number=state.item.item_number,
                rule_order=rule.order if rule else None,
                rule_description=rule.description if rule else "Sem regra cadastrada",
                competence=competence,
                index_reference=index_reference,
                is_business_day=is_business_day(day, holidays),
                opening_principal=data["opening"]["principal"],
                opening_correction=data["opening"]["correction"],
                opening_interest=data["opening"]["interest"],
                opening_penalty=data["opening"]["penalty"],
                correction_rate=correction_rate,
                correction_amount=data["correction_amount"],
                interest_rate=data["interest_rate"],
                interest_amount=data["interest_amount"],
                penalty_rate=data["penalty_rate"],
                penalty_amount=data["penalty_amount"],
                abatement_amount=applied,
                abatement_principal=ab_p,
                abatement_correction=ab_c,
                abatement_interest=ab_i,
                abatement_penalty=ab_m,
                closing_principal=state.principal,
                closing_correction=state.correction,
                closing_interest=state.interest,
                closing_penalty=state.penalty,
                closing_total=state.total,
            ))
        day += timedelta(days=1)

    # Se a apuração for interrompida por ausência de um fator futuro, fechamos
    # os juros acumulados até a última data efetivamente processada. A ausência
    # do fator seguinte não pode zerar juros já determináveis até a data válida.
    if not is_complete and daily_rows and max(r.day for r in daily_rows) < base_date:
        processed_end = max(r.day for r in daily_rows)
        for rule in rules_list:
            end_applied = (
                rule.interest_method in {INTEREST_SIMPLE_360_TJRJ, INTEREST_SIMPLE_YEARFRAC, INTEREST_LEGAL_BCB}
                or (rule.interest_application == INTEREST_APPLICATION_END and rule.interest_method in {INTEREST_SIMPLE_30, INTEREST_SIMPLE_COMPETENCE})
            )
            if not end_applied or (rule.monthly_interest_rate == 0 and rule.interest_method != INTEREST_LEGAL_BCB):
                continue
            for state in states:
                if state.item.origin_date >= processed_end or state.total_interest_generated != 0:
                    continue
                terminal = min(processed_end, rule.end_date)
                if rule.interest_method in {INTEREST_SIMPLE_360_TJRJ, INTEREST_SIMPLE_YEARFRAC, INTEREST_LEGAL_BCB}:
                    effective_start = max(rule.interest_start_date or rule.start_date, state.item.origin_date)
                    if terminal < effective_start:
                        continue
                    if rule.interest_method == INTEREST_SIMPLE_360_TJRJ:
                        rate = tjrj_simple_360_rate(rule.monthly_interest_rate, effective_start, terminal)
                    elif rule.interest_method == INTEREST_SIMPLE_YEARFRAC:
                        rate = yearfrac_simple_rate(rule.monthly_interest_rate, effective_start, terminal, rule.yearfrac_basis)
                    else:
                        legal_series = indices.get("TAXA_LEGAL")
                        if legal_series is None:
                            raise MissingIndexError(terminal, "TAXA_LEGAL", f"{terminal.year:04d}-{terminal.month:02d}")
                        rate = legal_rate_simple_accumulated(legal_series, effective_start, terminal)
                else:
                    key = (state.item.item_number, rule.order)
                    rate = deferred_interest_rates.get(key, Decimal("0"))
                amount = _apply_interest(state, rate, rule, round_daily_money, money_places)
                if amount == 0:
                    continue
                # Reconciliar a última linha diária do item com o fechamento dos juros.
                for idx in range(len(daily_rows) - 1, -1, -1):
                    row = daily_rows[idx]
                    if row.item_number == state.item.item_number and row.day == terminal:
                        daily_rows[idx] = replace(
                            row, interest_rate=rate, interest_amount=row.interest_amount + amount,
                            closing_interest=state.interest, closing_total=state.total
                        )
                        break

    if not is_complete and stopped_at is not None and daily_rows and max(r.day for r in daily_rows) >= base_date:
        interest_active = any(
            r.interest_method != INTEREST_NONE and (r.monthly_interest_rate != 0 or r.interest_method == INTEREST_LEGAL_BCB)
            for r in rules_list
        )
        if interest_active:
            warnings.append(
                f"A correção monetária ficou limitada até {(stopped_at - timedelta(days=1)).strftime('%d/%m/%Y')}; "
                f"os juros configurados foram apurados até a data-base de {base_date.strftime('%d/%m/%Y')}, "
                "sobre a base monetária disponível, sem projeção do índice ausente."
            )
        else:
            warnings.append(
                f"A correção monetária ficou limitada até {(stopped_at - timedelta(days=1)).strftime('%d/%m/%Y')}. "
                "Não foram aplicados juros porque a opção de juros está desativada."
            )

    if missing_gap_dates:
        first = min(missing_gap_dates)
        last = max(missing_gap_dates)
        warnings.append(
            "Há datas sem regra de atualização cadastrada entre "
            f"{first.strftime('%d/%m/%Y')} e {last.strftime('%d/%m/%Y')}. "
            "Nesses dias não foram aplicados correção nem juros."
        )

    summaries = [
        UpdateItemSummary(
            item_number=state.item.item_number,
            description=state.item.description,
            origin_date=state.item.origin_date,
            original_amount=state.item.amount,
            correction_amount=state.total_correction_generated,
            interest_amount=state.total_interest_generated,
            penalty_amount=state.total_penalty_generated,
            abatements=state.total_abatements,
            updated_total=state.total,
        )
        for state in states
    ]
    return MonetaryUpdateResult(
        base_date=base_date,
        items=summaries,
        daily_rows=daily_rows,
        rules=rules_list,
        abatements=abatements_list,
        penalties=penalties_list,
        warnings=warnings,
        is_complete=is_complete,
        stopped_at=stopped_at,
        missing_index_code=missing_code,
        missing_index_reference=missing_ref,
    )


def _style_sheet(ws, header_row: int = 1) -> None:
    for cell in ws[header_row]:
        cell.fill = PatternFill("solid", fgColor=PETROLEUM)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.border = Border(bottom=THIN)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        width = 12
        for cell in ws[get_column_letter(col)]:
            width = min(max(width, len(str(cell.value or "")) + 2), 34)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = ws.dimensions


def build_monetary_update_workbook(result: MonetaryUpdateResult, money_places: int = 2, percent_places: int = 4) -> bytes:
    money_fmt = 'R$ #,##0.' + ('0' * money_places) if money_places else 'R$ #,##0'
    pct_fmt = '0.' + ('0' * percent_places) + '%' if percent_places else '0%'
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws["A1"] = "ATUALIZAÇÃO MONETÁRIA DO SALDO DEVEDOR"
    ws["A1"].fill = PatternFill("solid", fgColor=PETROLEUM)
    ws["A1"].font = Font(color=WHITE, bold=True, size=14)
    ws.merge_cells("A1:D1")
    summary_rows = [
        ("Data-base", result.base_date),
        ("Status", "Completo" if result.is_complete else "Parcial"),
        ("Valor original", float(result.original_total)),
        ("Correção monetária", float(result.correction_total)),
        ("Juros", float(result.interest_total)),
        ("Multa", float(result.penalty_total)),
        ("Abatimentos", float(result.abatements_total)),
        ("Saldo atualizado", float(result.updated_total)),
        ("Índice ausente", f"{result.missing_index_code or ''} {result.missing_index_reference or ''}".strip()),
    ]
    for idx, (label, value) in enumerate(summary_rows, start=3):
        ws.cell(idx, 1, label).font = Font(bold=True, color=PETROLEUM)
        ws.cell(idx, 2, value)
        if label == "Data-base":
            ws.cell(idx, 2).number_format = "dd/mm/yyyy"
        if label in {"Valor original", "Correção monetária", "Juros", "Multa", "Abatimentos", "Saldo atualizado"}:
            ws.cell(idx, 2).number_format = money_fmt
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 42

    items = wb.create_sheet("Valores_Atualizados")
    items.append(["Item", "Descrição", "Data de origem", "Valor original", "Correção", "Juros", "Multa", "Abatimentos", "Saldo atualizado"])
    for x in result.items:
        items.append([
            x.item_number, x.description, x.origin_date, float(x.original_amount), float(x.correction_amount),
            float(x.interest_amount), float(x.penalty_amount), float(x.abatements), float(x.updated_total),
        ])
    _style_sheet(items)
    for cell in items["C"][1:]: cell.number_format = "dd/mm/yyyy"
    for col in range(4, 10):
        for cell in items[get_column_letter(col)][1:]: cell.number_format = money_fmt

    rules = wb.create_sheet("Regras")
    rules.append([
        "Ordem", "Início", "Fim", "Descrição", "Índice", "Método de correção", "Defasagem",
        "Juros (% a.m.)", "Início dos juros", "Método de juros", "Aplicação dos juros", "Base FRAÇÃOANO", "Base dos juros", "Ordem de incidência", "Competência inicia em",
    ])
    for r in result.rules:
        rules.append([
            r.order, r.start_date, r.end_date, r.description, r.correction_index_code,
            r.correction_method, r.index_lag_months, float(r.monthly_interest_rate), r.interest_start_date, r.interest_method,
            r.interest_application, r.yearfrac_basis, r.interest_base, r.event_order, r.competence_start_day,
        ])
    _style_sheet(rules)
    for col in (2,3,9):
        for cell in rules[get_column_letter(col)][1:]: cell.number_format = "dd/mm/yyyy"
    for cell in rules["H"][1:]: cell.number_format = pct_fmt

    penalties_ws = wb.create_sheet("Multas")
    penalties_ws.append(["Data", "Percentual", "Base", "Item alvo (0=todos)", "Observação"])
    for x in result.penalties:
        penalties_ws.append([x.event_date, float(x.rate), x.base, x.target_item, x.note])
    _style_sheet(penalties_ws)
    for cell in penalties_ws["A"][1:]: cell.number_format = "dd/mm/yyyy"
    for cell in penalties_ws["B"][1:]: cell.number_format = pct_fmt

    ab = wb.create_sheet("Abatimentos")
    ab.append(["Data", "Valor", "Item alvo (0=todos)", "Observação"])
    for x in result.abatements:
        ab.append([x.event_date, float(x.amount), x.target_item, x.note])
    _style_sheet(ab)
    for cell in ab["A"][1:]: cell.number_format = "dd/mm/yyyy"
    for cell in ab["B"][1:]: cell.number_format = money_fmt

    daily = wb.create_sheet("Memoria_Diaria")
    daily.append([
        "Data", "Item", "Regra", "Descrição da regra", "Competência", "Índice referência", "Dia útil",
        "Principal inicial", "Correção acumulada inicial", "Juros acumulados iniciais", "Multa acumulada inicial",
        "Taxa correção", "Correção do dia", "Taxa juros", "Juros do dia", "Taxa multa", "Multa do dia",
        "Abatimento total", "Abat. principal", "Abat. correção", "Abat. juros", "Abat. multa",
        "Principal final", "Correção final", "Juros final", "Multa final", "Saldo final",
    ])
    for x in result.daily_rows:
        daily.append([
            x.day, x.item_number, x.rule_order, x.rule_description, x.competence, x.index_reference,
            "Sim" if x.is_business_day else "Não", float(x.opening_principal), float(x.opening_correction),
            float(x.opening_interest), float(x.opening_penalty), float(x.correction_rate), float(x.correction_amount), float(x.interest_rate),
            float(x.interest_amount), float(x.penalty_rate), float(x.penalty_amount), float(x.abatement_amount), float(x.abatement_principal),
            float(x.abatement_correction), float(x.abatement_interest), float(x.abatement_penalty), float(x.closing_principal),
            float(x.closing_correction), float(x.closing_interest), float(x.closing_penalty), float(x.closing_total),
        ])
    _style_sheet(daily)
    for cell in daily["A"][1:]: cell.number_format = "dd/mm/yyyy"
    for col in (12, 14, 16):
        for cell in daily[get_column_letter(col)][1:]: cell.number_format = pct_fmt
    for col in list(range(8,12)) + [13,15,17] + list(range(18,27)):
        for cell in daily[get_column_letter(col)][1:]: cell.number_format = money_fmt

    warnings = wb.create_sheet("Avisos")
    warnings.append(["Aviso"])
    for warning in result.warnings or ["Nenhum aviso gerado."]:
        warnings.append([warning])
    _style_sheet(warnings)
    warnings["A1"].fill = PatternFill("solid", fgColor=ORANGE)
    warnings.column_dimensions["A"].width = 110

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def build_import_template() -> bytes:
    wb = Workbook()
    values = wb.active
    values.title = "Valores"
    values.append(["Prestação", "Data de origem", "Valor", "Descrição"])
    values.append([1, date(2024, 1, 20), 1000.00, "Exemplo - substitua ou exclua"])
    _style_sheet(values)
    values["B2"].number_format = "dd/mm/yyyy"
    values["C2"].number_format = 'R$ #,##0.00'

    ab = wb.create_sheet("Abatimentos")
    ab.append(["Data", "Valor", "Prestação alvo", "Observação"])
    ab.append([date(2024, 6, 10), 100.00, 0, "0 = distribuir entre todas as prestações abertas"])
    _style_sheet(ab)
    ab["A2"].number_format = "dd/mm/yyyy"
    ab["B2"].number_format = 'R$ #,##0.00'

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
