from __future__ import annotations

from dataclasses import dataclass, field, replace
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
import re
import unicodedata
from typing import BinaryIO, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo

from core.indices import IndexSeries


PERCENT = "PERCENTUAL"
FACTOR = "FATOR"
ECONOMIC = "ECONÔMICO"
JUDICIAL = "JUDICIAL"

_REQUIRED_CATALOG = {"codigo", "nome", "categoria", "tipovalor", "fonte", "status", "ativo"}
_REQUIRED_SERIES = {"codigo", "competencia", "valor"}
_COMPETENCE_RE = re.compile(r"^(\d{4})-(0[1-9]|1[0-2])$")


@dataclass(frozen=True)
class FactorSeriesView:
    code: str
    name: str
    values: dict[str, Decimal]
    source: str = ""
    observation: str = ""

    @property
    def coverage_start(self) -> str:
        return min(self.values) if self.values else "-"

    @property
    def coverage_end(self) -> str:
        return max(self.values) if self.values else "-"


@dataclass
class SeriesDefinition:
    code: str
    name: str
    category: str
    value_type: str
    periodicity: str = "Mensal"
    source: str = ""
    status: str = ""
    active: bool = True
    observation: str = ""
    # Valores no MESMO formato apresentado na planilha:
    # PERCENTUAL => pontos percentuais (0,35 significa 0,35%); FATOR => fator bruto.
    values: dict[str, Decimal] = field(default_factory=dict)
    row_metadata: dict[str, dict[str, str]] = field(default_factory=dict)

    @property
    def first_competence(self) -> str:
        return min(self.values) if self.values else "-"

    @property
    def last_competence(self) -> str:
        return max(self.values) if self.values else "-"

    @property
    def record_count(self) -> int:
        return len(self.values)

    def as_index_series(self) -> IndexSeries:
        if self.value_type != PERCENT:
            raise ValueError(f"{self.code} não é uma série percentual.")
        converted = {k: (v / Decimal("100")) for k, v in self.values.items()}
        return IndexSeries(code=self.code, values=converted, source=self.source)

    def as_factor_series(self) -> FactorSeriesView:
        if self.value_type != FACTOR:
            raise ValueError(f"{self.code} não é uma série de fatores.")
        return FactorSeriesView(
            code=self.code,
            name=self.name,
            values=dict(self.values),
            source=self.source,
            observation=self.observation,
        )


@dataclass
class IndexRepository:
    definitions: dict[str, SeriesDefinition]
    source_path: str = ""
    warnings: tuple[str, ...] = tuple()

    @property
    def active_definitions(self) -> dict[str, SeriesDefinition]:
        return {k: v for k, v in self.definitions.items() if v.active}

    @property
    def percent_series(self) -> dict[str, IndexSeries]:
        return {
            code: item.as_index_series()
            for code, item in self.active_definitions.items()
            if item.value_type == PERCENT
        }

    @property
    def factor_series(self) -> dict[str, FactorSeriesView]:
        return {
            code: item.as_factor_series()
            for code, item in self.active_definitions.items()
            if item.value_type == FACTOR
        }

    def get(self, code: str) -> SeriesDefinition:
        key = normalize_code(code)
        if key not in self.definitions:
            raise KeyError(f"Série {key} não encontrada na base única de índices.")
        return self.definitions[key]

    def status_rows(self) -> list[dict[str, object]]:
        rows = []
        for code, item in sorted(self.definitions.items()):
            rows.append({
                "Código": code,
                "Nome": item.name,
                "Categoria": item.category,
                "Tipo": "Percentual" if item.value_type == PERCENT else "Fator",
                "Primeira competência": item.first_competence,
                "Última competência": item.last_competence,
                "Registros": item.record_count,
                "Status": item.status or "-",
                "Ativo": "Sim" if item.active else "Não",
                "Fonte": item.source or "-",
            })
        return rows


@dataclass(frozen=True)
class IndexDiscovery:
    repository: IndexRepository | None
    path: Path
    error: str = ""


def _norm(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def normalize_code(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or "").strip())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")
    return text.upper()


def _as_bool(value: object) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return _norm(value) in {"sim", "s", "true", "verdadeiro", "1", "ativo", "ativa", "x"}


def _as_decimal(value: object, *, label: str) -> Decimal:
    if isinstance(value, Decimal):
        return value
    if value is None or str(value).strip() == "":
        raise ValueError(f"{label}: valor vazio.")
    text = str(value).strip().replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", ".")
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"{label}: valor numérico inválido ({value}).") from exc


def _open_workbook(source: str | Path | bytes | BinaryIO):
    if isinstance(source, (str, Path)):
        return load_workbook(source, data_only=True, read_only=False)
    if isinstance(source, bytes):
        return load_workbook(BytesIO(source), data_only=True, read_only=False)
    return load_workbook(source, data_only=True, read_only=False)


def _header_map(ws) -> dict[str, int]:
    return {_norm(cell.value): idx for idx, cell in enumerate(ws[1]) if cell.value is not None}


def _require_headers(ws, required: set[str], sheet_name: str) -> dict[str, int]:
    header = _header_map(ws)
    missing = sorted(required - set(header))
    if missing:
        raise ValueError(f"A aba {sheet_name} não possui as colunas obrigatórias: {', '.join(missing)}.")
    return header


def _value(row: tuple, header: dict[str, int], *names: str, default=None):
    for name in names:
        key = _norm(name)
        if key in header and header[key] < len(row):
            value = row[header[key]]
            if value is not None:
                return value
    return default


def load_index_workbook(source: str | Path | bytes | BinaryIO) -> IndexRepository:
    wb = _open_workbook(source)
    if "CATALOGO" not in wb.sheetnames or "SERIES" not in wb.sheetnames:
        raise ValueError("A base de índices deve conter as abas CATALOGO e SERIES.")

    catalog_ws = wb["CATALOGO"]
    series_ws = wb["SERIES"]
    ch = _require_headers(catalog_ws, _REQUIRED_CATALOG, "CATALOGO")
    sh = _require_headers(series_ws, _REQUIRED_SERIES, "SERIES")

    definitions: dict[str, SeriesDefinition] = {}
    for row_number, row in enumerate(catalog_ws.iter_rows(min_row=2, values_only=True), start=2):
        code = normalize_code(_value(row, ch, "Código"))
        if not code:
            continue
        if code in definitions:
            raise ValueError(f"CATALOGO: código duplicado {code} (linha {row_number}).")
        value_type = normalize_code(_value(row, ch, "Tipo_Valor", "Tipo Valor"))
        if value_type not in {PERCENT, FACTOR}:
            raise ValueError(f"CATALOGO: Tipo_Valor inválido para {code}. Use PERCENTUAL ou FATOR.")
        category_raw = str(_value(row, ch, "Categoria", default=ECONOMIC) or ECONOMIC).strip().upper()
        category = JUDICIAL if _norm(category_raw) == "judicial" else ECONOMIC
        definitions[code] = SeriesDefinition(
            code=code,
            name=str(_value(row, ch, "Nome", "Nome exibido", default=code) or code).strip(),
            category=category,
            value_type=value_type,
            periodicity=str(_value(row, ch, "Periodicidade", default="Mensal") or "Mensal").strip(),
            source=str(_value(row, ch, "Fonte", default="") or "").strip(),
            status=str(_value(row, ch, "Status", "Situação", default="") or "").strip(),
            active=_as_bool(_value(row, ch, "Ativo", default=True)),
            observation=str(_value(row, ch, "Observação", "Observacao", default="") or "").strip(),
        )

    if not definitions:
        raise ValueError("A aba CATALOGO não contém séries cadastradas.")

    duplicates: set[tuple[str, str]] = set()
    for row_number, row in enumerate(series_ws.iter_rows(min_row=2, values_only=True), start=2):
        code = normalize_code(_value(row, sh, "Código"))
        competence = str(_value(row, sh, "Competência", "Competencia", default="") or "").strip()
        raw_value = _value(row, sh, "Valor")
        if not code and not competence and raw_value in (None, ""):
            continue
        if code not in definitions:
            raise ValueError(f"SERIES: código {code or '-'} não existe no CATALOGO (linha {row_number}).")
        if not _COMPETENCE_RE.fullmatch(competence):
            raise ValueError(f"SERIES: competência inválida '{competence}' na linha {row_number}. Use AAAA-MM.")
        pair = (code, competence)
        if pair in duplicates:
            raise ValueError(f"SERIES: competência duplicada para {code}: {competence}.")
        duplicates.add(pair)

        item = definitions[code]
        row_type_raw = _value(row, sh, "Tipo_Valor", "Tipo Valor", default=item.value_type)
        row_type = normalize_code(row_type_raw)
        if row_type and row_type != item.value_type:
            raise ValueError(
                f"SERIES: Tipo_Valor de {code}/{competence} ({row_type}) diverge do CATALOGO ({item.value_type})."
            )
        value = _as_decimal(raw_value, label=f"{code}/{competence}")
        if item.value_type == FACTOR and value == 0:
            raise ValueError(f"SERIES: fator zero não é permitido ({code}/{competence}).")
        item.values[competence] = value
        item.row_metadata[competence] = {
            "source": str(_value(row, sh, "Fonte", default=item.source) or item.source).strip(),
            "status": str(_value(row, sh, "Status", default=item.status) or item.status).strip(),
            "active": "SIM" if _as_bool(_value(row, sh, "Ativo", default=item.active)) else "NÃO",
            "observation": str(_value(row, sh, "Observação", "Observacao", default="") or "").strip(),
        }

    warnings: list[str] = []
    for code, item in definitions.items():
        if item.active and not item.values:
            warnings.append(f"{code}: série ativa sem valores.")

    source_path = str(source) if isinstance(source, (str, Path)) else "Base carregada em memória"
    return IndexRepository(definitions=definitions, source_path=source_path, warnings=tuple(warnings))


def _candidate_files(app_dir: Path, preferred_path: Path | None) -> list[Path]:
    config_dir = app_dir / "config"
    known = [
        "indices_motor_calculos.xlsx",
        "Indices_Motor_Calculos.xlsx",
        "base_indices_motor_calculos.xlsx",
        "Base_Indices_Calculos_v1.xlsx",
    ]
    ordered: list[Path] = []
    if preferred_path:
        ordered.append(preferred_path)
    for folder in (config_dir, app_dir):
        ordered.extend(folder / name for name in known)
        if folder.exists():
            ordered.extend(
                p for p in sorted(folder.glob("*.xlsx"))
                if not p.name.startswith("~$") and "indice" in _norm(p.stem)
            )
    seen: set[str] = set()
    result: list[Path] = []
    for path in ordered:
        key = str(path.resolve()) if path.exists() else str(path.absolute())
        if key not in seen:
            seen.add(key)
            result.append(path)
    return result


def discover_index_workbook(app_dir: str | Path, preferred_path: str | Path | None = None) -> IndexDiscovery:
    app_dir = Path(app_dir)
    preferred = Path(preferred_path) if preferred_path else None
    valid: list[tuple[tuple[int, int], int, Path, IndexRepository]] = []
    errors: list[str] = []
    for order, path in enumerate(_candidate_files(app_dir, preferred)):
        if not path.exists() or not path.is_file():
            continue
        try:
            repo = load_index_workbook(path)
        except Exception as exc:
            errors.append(f"{path.name}: {exc}")
            continue
        score = (len(repo.active_definitions), sum(x.record_count for x in repo.active_definitions.values()))
        valid.append((score, -order, path, repo))
    if valid:
        valid.sort(key=lambda x: (x[0], x[1]), reverse=True)
        _, _, path, repo = valid[0]
        repo.source_path = str(path)
        return IndexDiscovery(repository=repo, path=path)
    fallback = preferred or (app_dir / "config" / "indices_motor_calculos.xlsx")
    return IndexDiscovery(repository=None, path=fallback, error="; ".join(errors[-3:]))


def clone_repository(repo: IndexRepository) -> IndexRepository:
    definitions: dict[str, SeriesDefinition] = {}
    for code, item in repo.definitions.items():
        definitions[code] = SeriesDefinition(
            code=item.code,
            name=item.name,
            category=item.category,
            value_type=item.value_type,
            periodicity=item.periodicity,
            source=item.source,
            status=item.status,
            active=item.active,
            observation=item.observation,
            values=dict(item.values),
            row_metadata={k: dict(v) for k, v in item.row_metadata.items()},
        )
    return IndexRepository(definitions=definitions, source_path=repo.source_path, warnings=repo.warnings)


def replace_series_values(
    repo: IndexRepository,
    code: str,
    values: dict[str, Decimal],
    *,
    source: str | None = None,
) -> IndexRepository:
    updated = clone_repository(repo)
    item = updated.get(code)
    validated: dict[str, Decimal] = {}
    for competence, value in values.items():
        if not _COMPETENCE_RE.fullmatch(str(competence)):
            raise ValueError(f"Competência inválida: {competence}. Use AAAA-MM.")
        dec = _as_decimal(value, label=f"{item.code}/{competence}")
        if item.value_type == FACTOR and dec == 0:
            raise ValueError(f"Fator zero não é permitido ({item.code}/{competence}).")
        validated[str(competence)] = dec
    if not validated:
        raise ValueError("A série precisa ter ao menos uma competência.")
    item.values = dict(sorted(validated.items()))
    if source:
        item.source = source.strip()
    # Preserva metadados das competências já existentes e inicializa as novas.
    item.row_metadata = {
        comp: dict(item.row_metadata.get(comp, {
            "source": item.source,
            "status": item.status,
            "active": "SIM" if item.active else "NÃO",
            "observation": "",
        }))
        for comp in item.values
    }
    return updated


def upsert_series(
    repo: IndexRepository,
    *,
    code: str,
    name: str,
    category: str,
    value_type: str,
    periodicity: str,
    source: str,
    status: str,
    active: bool,
    observation: str,
    values: dict[str, Decimal],
) -> IndexRepository:
    key = normalize_code(code)
    if not key:
        raise ValueError("Informe o código da série.")
    vtype = normalize_code(value_type)
    if vtype not in {PERCENT, FACTOR}:
        raise ValueError("Tipo de valor inválido. Use PERCENTUAL ou FATOR.")
    category_norm = JUDICIAL if _norm(category) == "judicial" else ECONOMIC
    updated = clone_repository(repo)
    updated.definitions[key] = SeriesDefinition(
        code=key,
        name=(name or key).strip(),
        category=category_norm,
        value_type=vtype,
        periodicity=(periodicity or "Mensal").strip(),
        source=(source or "Criado na sessão").strip(),
        status=(status or "EM VALIDAÇÃO").strip(),
        active=bool(active),
        observation=(observation or "").strip(),
    )
    return replace_series_values(updated, key, values, source=source)


def repository_to_xlsx_bytes(repo: IndexRepository) -> bytes:
    """Gera a base única padronizada para download/redeploy.

    Essa rotina pertence ao aplicativo. Ela usa openpyxl porque será executada no
    Databricks App; não altera a matemática dos motores.
    """
    wb = Workbook()
    ws_readme = wb.active
    ws_readme.title = "LEIA-ME"
    ws_cat = wb.create_sheet("CATALOGO")
    ws_series = wb.create_sheet("SERIES")
    ws_control = wb.create_sheet("CONTROLE")

    navy = "073D6D"
    petroleum = "005A70"
    orange = "F47C20"
    light = "EEF5F7"
    border = Side(style="thin", color="D6E2E5")

    readme = [
        ("BASE ÚNICA DE ÍNDICES — MOTOR DE CÁLCULOS", ""),
        ("Arquivo oficial", "config/indices_motor_calculos.xlsx"),
        ("Manutenção", "Inclua/edite competências na aba SERIES e importe a base pela Central de índices."),
        ("Percentuais", "0,35 na coluna Valor representa 0,35% no mês."),
        ("Fatores", "Fatores judiciais permanecem brutos; a regra de negócio define como utilizá-los."),
        ("Persistência", "Alterações feitas pela interface valem na sessão até a base exportada ser colocada no deploy."),
    ]
    for row in readme:
        ws_readme.append(row)

    cat_headers = ["Código", "Nome", "Categoria", "Tipo_Valor", "Periodicidade", "Fonte", "Status", "Ativo", "Observação"]
    ws_cat.append(cat_headers)
    series_headers = ["Código", "Nome", "Categoria", "Competência", "Valor", "Tipo_Valor", "Fonte", "Status", "Ativo", "Observação"]
    ws_series.append(series_headers)
    control_headers = ["Código", "Nome", "Categoria", "Tipo_Valor", "Primeira competência", "Última competência", "Registros", "Status", "Fonte", "Situação"]
    ws_control.append(control_headers)

    for code, item in sorted(repo.definitions.items()):
        ws_cat.append([
            code, item.name, item.category, item.value_type, item.periodicity,
            item.source, item.status, "SIM" if item.active else "NÃO", item.observation,
        ])
        situation = "ATENÇÃO" if any(token in item.status.upper() for token in ("VALIDAR", "PARCIAL", "REVIS")) else "DISPONÍVEL"
        ws_control.append([
            code, item.name, item.category, item.value_type, item.first_competence,
            item.last_competence, item.record_count, item.status, item.source, situation,
        ])
        for comp, value in sorted(item.values.items()):
            meta = item.row_metadata.get(comp, {})
            ws_series.append([
                code, item.name, item.category, comp, float(value), item.value_type,
                meta.get("source") or item.source,
                meta.get("status") or item.status,
                meta.get("active") or ("SIM" if item.active else "NÃO"),
                meta.get("observation") or "",
            ])

    for ws in (ws_cat, ws_series, ws_control):
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor=navy)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows():
            for cell in row:
                cell.border = Border(left=border, right=border, top=border, bottom=border)
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    for ws, widths in (
        (ws_cat, [20, 32, 14, 16, 15, 52, 28, 9, 70]),
        (ws_series, [20, 32, 14, 14, 18, 16, 52, 28, 9, 70]),
        (ws_control, [20, 32, 14, 16, 18, 18, 12, 28, 52, 15]),
    ):
        for idx, width in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + idx)].width = width

    if ws_cat.max_row >= 2:
        table = Table(displayName="CatalogoIndices", ref=f"A1:I{ws_cat.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        ws_cat.add_table(table)
    if ws_series.max_row >= 2:
        table = Table(displayName="SeriesIndices", ref=f"A1:J{ws_series.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        ws_series.add_table(table)
    if ws_control.max_row >= 2:
        table = Table(displayName="ControleIndices", ref=f"A1:J{ws_control.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True, showFirstColumn=False, showLastColumn=False)
        ws_control.add_table(table)

    ws_readme.merge_cells("A1:B1")
    ws_readme["A1"].fill = PatternFill("solid", fgColor=navy)
    ws_readme["A1"].font = Font(bold=True, color="FFFFFF", size=14)
    ws_readme.column_dimensions["A"].width = 22
    ws_readme.column_dimensions["B"].width = 105
    for row in ws_readme.iter_rows(min_row=2):
        row[0].fill = PatternFill("solid", fgColor=light)
        row[0].font = Font(bold=True, color=petroleum)
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=border, right=border, top=border, bottom=border)

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
